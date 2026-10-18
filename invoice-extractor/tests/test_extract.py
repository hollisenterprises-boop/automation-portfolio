import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import extract as ex  # noqa: E402
import make_sample_invoices as samples  # noqa: E402

AS_OF = date(2026, 10, 18)


@pytest.fixture(scope="module")
def folder(tmp_path_factory):
    out = tmp_path_factory.mktemp("invoices")
    for inv in samples.sample_invoices():
        samples.write_invoice(inv, out)
    return out


@pytest.fixture(scope="module")
def invoices(folder):
    return {i.file: i for i in ex.read_folder(folder)}


@pytest.fixture(scope="module")
def expected():
    return {i.file_name: i for i in samples.sample_invoices()}


def test_reads_every_file(invoices, expected):
    assert set(invoices) == set(expected)


@pytest.mark.parametrize("layout", ["table", "lines", "services"])
def test_reads_each_layout(invoices, expected, layout):
    for name, want in expected.items():
        if want.layout != layout:
            continue
        got = invoices[name]
        assert got.number == want.number
        assert got.issued == want.issued
        assert got.due == want.due
        assert got.subtotal == pytest.approx(want.subtotal)
        assert got.gst == pytest.approx(want.gst)
        assert got.total == pytest.approx(want.total)
        assert len(got.lines) == len(want.lines)
        assert [li.amount for li in got.lines] == pytest.approx([round(q * p, 2) for _, q, p in want.lines])


def test_supplier_names(invoices):
    assert {i.supplier for i in invoices.values()} == {
        "Coastal Packaging Supplies", "Northside Print Co.", "Bright Spark Electrical",
    }


def test_flags_total_that_does_not_add_up(invoices):
    problems = invoices["coastal_INV-10512.pdf"].problems
    assert len(problems) == 1
    assert "but the total says $653.17 (+$10.00)" in problems[0]


def test_flags_wrong_gst(invoices):
    problems = invoices["brightspark_BSE-0102.pdf"].problems
    assert problems == ["GST charged is $99.69, but 10% of the subtotal is $79.75"]


def test_flags_duplicate_and_leaves_it_out(invoices):
    dup = invoices["scan_0921_reprint.pdf"]
    assert dup.duplicate_of == "northside_NP-2344.pdf"
    assert invoices["northside_NP-2344.pdf"].duplicate_of is None
    assert dup not in ex.counted(list(invoices.values()))


def test_clean_invoices_have_no_problems(invoices):
    flagged = {name for name, inv in invoices.items() if inv.problems}
    assert flagged == {"coastal_INV-10512.pdf", "brightspark_BSE-0102.pdf", "scan_0921_reprint.pdf"}


def test_due_date_from_payment_terms(invoices):
    assert invoices["brightspark_BSE-0119.pdf"].due == date(2026, 10, 10)  # 3 Oct + 7 days


@pytest.mark.parametrize("text, want", [
    ("04/08/2026", date(2026, 8, 4)),  # day first
    ("2026-08-04", date(2026, 8, 4)),
    ("4 Aug 2026", date(2026, 8, 4)),
    ("4 August 2026", date(2026, 8, 4)),
    ("August 4 2026", date(2026, 8, 4)),
    ("next Tuesday", None),
])
def test_parse_date(text, want):
    assert ex.parse_date(text) == want


def test_missing_fields_are_listed_not_guessed():
    inv = ex.parse_text("Some Supplier\nThanks for your order\n", "odd.pdf")
    assert inv.number is None and inv.total is None
    assert "No invoice number found" in inv.problems
    assert "No total found" in inv.problems
    assert "No line items found" in inv.problems


def test_line_that_does_not_multiply_out():
    inv = ex.Invoice(file="x.pdf", lines=[ex.LineItem("Widget", 3, 2.50, 8.00)])
    ex.check_maths(inv)
    assert inv.problems == ['Line "Widget": 3 x $2.50 = $7.50, but the invoice says $8.00']


def test_due_status():
    inv = ex.Invoice(file="x.pdf", due=date(2026, 10, 25))
    assert ex.due_status(inv, AS_OF) == "Due soon"
    inv.due = date(2026, 10, 1)
    assert ex.due_status(inv, AS_OF) == "Past due"
    inv.due = date(2026, 12, 1)
    assert ex.due_status(inv, AS_OF) == "Later"


def test_workbook(invoices, tmp_path):
    path = tmp_path / "out.xlsx"
    ex.build_workbook(list(invoices.values()), AS_OF).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Invoices", "Line Items", "Checks"]
    assert wb["Invoices"].max_row == 1 + 13
    assert wb["Checks"].max_row == 1 + 3
    real = ex.counted(list(invoices.values()))
    assert wb["Summary"]["B4"].value == 12
    assert wb["Summary"]["B5"].value == pytest.approx(round(sum(i.total for i in real), 2))
