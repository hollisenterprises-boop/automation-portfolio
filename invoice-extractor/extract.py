"""Read a folder of supplier invoice PDFs and write one checked Excel workbook.

    python extract.py sample_data/invoices -o output/invoices.xlsx --as-of 2026-10-18

For each PDF it finds the supplier, invoice number, dates, line items,
subtotal, GST and total, whichever of the common layouts the supplier uses.
Then it checks the maths (each line, the items against the subtotal, subtotal
plus GST against the total, GST at 10%), flags the same invoice saved twice,
and shows what falls due in the next two weeks. Anything it can't read is
listed for a person to check rather than guessed.
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

GST_RATE = 0.10
DUE_SOON_DAYS = 14

# Slashed dates are read day-first (Australian format).
DATE_FORMATS = ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%b %d %Y", "%B %d %Y")

AMOUNT = r"\$?\s*([\d,]+\.\d{2})"
FIELDS = {
    "number": [r"^Invoice\s*(?:No|Number|#)\.?\s*:?\s*([A-Z0-9][\w/-]*)", r"^Tax Invoice\s*#\s*([A-Z0-9][\w/-]*)"],
    "issued": [r"^Invoice Date\s*:\s*(.+)$", r"^Issued\s*:\s*(.+)$", r"^Date\s*:\s*(.+)$"],
    "due": [r"^Due Date\s*:\s*(.+)$", r"^Payment due\s*:\s*(.+)$", r"^Due\s*:\s*(.+)$"],
    "terms": [r"^Terms\s*:\s*(\d+)\s*days"],
    "subtotal": [rf"^Sub-?total(?:\s*\(ex GST\))?\s*:\s*{AMOUNT}", rf"^Total ex GST\s*:\s*{AMOUNT}"],
    "gst": [rf"^GST(?:\s*\(?10%\)?)?\s*:\s*{AMOUNT}"],
    "total": [rf"^Total(?: due)?\s*:\s*{AMOUNT}", rf"^Amount (?:Payable|Due)\s*:\s*{AMOUNT}", rf"^Balance due\s*:\s*{AMOUNT}"],
}
# "Shipping boxes 40x30x20cm 200 $1.15 $230.00" (ruled tables, read as text)
TABLE_ROW = re.compile(r"^(?P<desc>.+?)\s+(?P<qty>\d+(?:\.\d+)?)\s+\$(?P<unit>[\d,]+\.\d{2,4})\s+\$(?P<amount>[\d,]+\.\d{2})$")
# "500 x Business cards @ $0.09 = $45.00"
X_AT_ROW = re.compile(r"^(?P<qty>\d+(?:\.\d+)?)\s*x\s+(?P<desc>.+?)\s+@\s+\$(?P<unit>[\d,]+\.\d{2,4})\s*=\s*\$(?P<amount>[\d,]+\.\d{2})$")
NOT_SUPPLIER = re.compile(r"^(tax\s+)?invoice\b", re.I)

MONEY = '"$"#,##0.00'
HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CHECK_FILL = PatternFill("solid", fgColor="FDE2E1")


@dataclass
class LineItem:
    description: str
    quantity: float
    unit_price: float
    amount: float


@dataclass
class Invoice:
    file: str
    supplier: str | None = None
    number: str | None = None
    issued: date | None = None
    due: date | None = None
    subtotal: float | None = None
    gst: float | None = None
    total: float | None = None
    lines: list[LineItem] = field(default_factory=list)
    problems: list[str] = field(default_factory=list)
    duplicate_of: str | None = None

    @property
    def ok(self) -> bool:
        return not self.problems


def to_money(text: str) -> float:
    return float(text.replace("$", "").replace(",", "").strip())


def parse_date(text: str) -> date | None:
    text = text.strip().rstrip(".")
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def find(text: str, name: str) -> str | None:
    for pattern in FIELDS[name]:
        match = re.search(pattern, text, re.I | re.M)
        if match:
            return match.group(1).strip()
    return None


def read_lines(text: str) -> list[LineItem]:
    items = []
    for raw in text.splitlines():
        line = raw.strip()
        match = TABLE_ROW.match(line) or X_AT_ROW.match(line)
        if match:
            items.append(LineItem(
                description=match["desc"].strip(),
                quantity=float(match["qty"]),
                unit_price=to_money(match["unit"]),
                amount=to_money(match["amount"]),
            ))
    return items


def parse_text(text: str, file_name: str) -> Invoice:
    """Pull the invoice fields out of a PDF's text and note anything missing."""
    inv = Invoice(file=file_name)
    lines = [l.strip() for l in text.splitlines() if l.strip()]
    inv.supplier = next((l for l in lines if not NOT_SUPPLIER.match(l)), None)
    inv.number = find(text, "number")

    issued = find(text, "issued")
    if issued:
        inv.issued = parse_date(issued)
        if inv.issued is None:
            inv.problems.append(f'Could not read the invoice date "{issued}"')
    else:
        inv.problems.append("No invoice date found")

    due = find(text, "due")
    terms = find(text, "terms")
    if due:
        inv.due = parse_date(due)
        if inv.due is None:
            inv.problems.append(f'Could not read the due date "{due}"')
    elif terms and inv.issued:
        inv.due = inv.issued + timedelta(days=int(terms))

    for name in ("subtotal", "gst", "total"):
        value = find(text, name)
        setattr(inv, name, to_money(value) if value else None)

    inv.lines = read_lines(text)
    if not inv.number:
        inv.problems.append("No invoice number found")
    if inv.total is None:
        inv.problems.append("No total found")
    if not inv.lines:
        inv.problems.append("No line items found")
    return inv


def check_maths(inv: Invoice) -> None:
    for item in inv.lines:
        expected = round(item.quantity * item.unit_price, 2)
        if abs(expected - item.amount) > max(0.01, item.quantity * 0.00005):
            inv.problems.append(
                f'Line "{item.description}": {item.quantity:g} x ${item.unit_price:,.2f} = ${expected:,.2f}, '
                f"but the invoice says ${item.amount:,.2f}"
            )
    if inv.lines and inv.subtotal is not None:
        items_total = round(sum(i.amount for i in inv.lines), 2)
        if abs(items_total - inv.subtotal) > 0.01:
            inv.problems.append(f"Line items add up to ${items_total:,.2f}, but the subtotal says ${inv.subtotal:,.2f}")
    if None not in (inv.subtotal, inv.gst, inv.total):
        expected_total = round(inv.subtotal + inv.gst, 2)
        diff = round(inv.total - expected_total, 2)
        if abs(diff) > 0.01:
            inv.problems.append(
                f"Subtotal ${inv.subtotal:,.2f} + GST ${inv.gst:,.2f} = ${expected_total:,.2f}, "
                f"but the total says ${inv.total:,.2f} ({'+' if diff > 0 else '-'}${abs(diff):,.2f})"
            )
    if inv.subtotal is not None and inv.gst is not None:
        expected_gst = round(inv.subtotal * GST_RATE, 2)
        if abs(expected_gst - inv.gst) > 0.02:
            inv.problems.append(f"GST charged is ${inv.gst:,.2f}, but 10% of the subtotal is ${expected_gst:,.2f}")
    if inv.issued and inv.due and inv.due < inv.issued:
        inv.problems.append("Due date is before the invoice date")


def mark_duplicates(invoices: list[Invoice]) -> None:
    """The same supplier and invoice number in two files is counted once."""
    seen: dict[tuple[str, str], str] = {}
    for inv in sorted(invoices, key=lambda i: i.file):
        if not (inv.supplier and inv.number):
            continue
        key = (inv.supplier.lower(), inv.number.upper())
        if key in seen:
            inv.duplicate_of = seen[key]
            inv.problems.append(f"Same invoice as {seen[key]} (possible duplicate, left out of the totals)")
        else:
            seen[key] = inv.file


def read_pdf(path: Path) -> Invoice:
    with pdfplumber.open(path) as pdf:
        text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    if not text.strip():
        inv = Invoice(file=path.name)
        inv.problems.append("No text in this PDF (probably a scan; needs OCR)")
        return inv
    inv = parse_text(text, path.name)
    check_maths(inv)
    return inv


def read_folder(folder: Path) -> list[Invoice]:
    invoices = [read_pdf(p) for p in sorted(folder.glob("*.pdf"))]
    mark_duplicates(invoices)
    return invoices


def due_status(inv: Invoice, as_of: date) -> str:
    if inv.due is None:
        return "Unknown"
    if inv.due < as_of:
        return "Past due"
    if inv.due <= as_of + timedelta(days=DUE_SOON_DAYS):
        return "Due soon"
    return "Later"


def counted(invoices: list[Invoice]) -> list[Invoice]:
    return [i for i in invoices if i.duplicate_of is None]


# ---------------------------------------------------------------- Excel output

def write_table(ws, headers: list[str], rows: list[list], money_cols: set[int] = frozenset(), widths=None) -> None:
    ws.append(headers)
    for cell in ws[ws.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    first = ws.max_row + 1
    for row in rows:
        ws.append(row)
    for r in range(first, ws.max_row + 1):
        for c in money_cols:
            ws.cell(r, c).number_format = MONEY
    for i, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = ws.cell(first, 1)
    ws.auto_filter.ref = f"A{first - 1}:{get_column_letter(len(headers))}{max(ws.max_row, first)}"


def build_workbook(invoices: list[Invoice], as_of: date) -> Workbook:
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    real = counted(invoices)
    soon = [i for i in real if due_status(i, as_of) == "Due soon"]
    flagged = [i for i in invoices if not i.ok]

    summary["A1"] = "Supplier invoices"
    summary["A1"].font = Font(bold=True, size=16)
    summary["A2"] = f"{len(invoices)} PDF files read, as of {as_of:%d %b %Y}"
    figures = [
        ("Invoices (duplicates removed)", len(real), None),
        ("Total spend, inc GST", sum(i.total or 0 for i in real), MONEY),
        ("GST to claim", sum(i.gst or 0 for i in real), MONEY),
        ("Invoices to check", len(flagged), None),
        (f"Due in the next {DUE_SOON_DAYS} days", sum(i.total or 0 for i in soon), MONEY),
    ]
    for row, (label, value, fmt) in enumerate(figures, start=4):
        summary.cell(row, 1, label)
        cell = summary.cell(row, 2, round(value, 2) if isinstance(value, float) else value)
        cell.font = Font(bold=True)
        if fmt:
            cell.number_format = fmt
    summary.cell(10, 1, "By supplier").font = Font(bold=True, size=12)
    by_supplier: dict[str, list[Invoice]] = defaultdict(list)
    for inv in real:
        by_supplier[inv.supplier or "Unknown"].append(inv)
    summary.append([])
    summary.append(["Supplier", "Invoices", "Ex GST", "GST", "Total"])
    for cell in summary[summary.max_row]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for name, items in sorted(by_supplier.items(), key=lambda kv: -sum(i.total or 0 for i in kv[1])):
        summary.append([name, len(items), sum(i.subtotal or 0 for i in items), sum(i.gst or 0 for i in items),
                        sum(i.total or 0 for i in items)])
        for c in (3, 4, 5):
            summary.cell(summary.max_row, c).number_format = MONEY
    for col, width in zip("ABCDE", (34, 14, 14, 12, 14)):
        summary.column_dimensions[col].width = width

    ws = wb.create_sheet("Invoices")
    write_table(
        ws,
        ["File", "Supplier", "Invoice No", "Invoice Date", "Due Date", "Due", "Subtotal", "GST", "Total", "Lines", "Status"],
        [[i.file, i.supplier, i.number, i.issued, i.due, due_status(i, as_of), i.subtotal, i.gst, i.total,
          len(i.lines), "OK" if i.ok else "Check"] for i in invoices],
        money_cols={7, 8, 9},
        widths=[28, 26, 13, 13, 13, 11, 12, 10, 12, 7, 9],
    )
    for r in range(2, ws.max_row + 1):
        for c in (4, 5):
            ws.cell(r, c).number_format = "dd/mm/yyyy"
        if ws.cell(r, 11).value == "Check":
            for c in range(1, 12):
                ws.cell(r, c).fill = CHECK_FILL

    write_table(
        wb.create_sheet("Line Items"),
        ["File", "Supplier", "Invoice No", "Description", "Qty", "Unit Price", "Amount"],
        [[i.file, i.supplier, i.number, li.description, li.quantity, li.unit_price, li.amount]
         for i in invoices for li in i.lines],
        money_cols={6, 7},
        widths=[28, 26, 13, 36, 8, 11, 12],
    )
    write_table(
        wb.create_sheet("Checks"),
        ["File", "Supplier", "Invoice No", "What to check"],
        [[i.file, i.supplier, i.number, p] for i in invoices for p in i.problems],
        widths=[28, 26, 13, 90],
    )
    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("folder", type=Path, help="folder of invoice PDFs")
    parser.add_argument("-o", "--output", type=Path, default=Path("output/invoices.xlsx"))
    parser.add_argument("--as-of", type=date.fromisoformat, default=date.today(), help="date for 'due soon' (YYYY-MM-DD)")
    args = parser.parse_args()

    invoices = read_folder(args.folder)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(invoices, args.as_of).save(args.output)
    real = counted(invoices)
    flagged = [i for i in invoices if not i.ok]
    print(f"{len(invoices)} PDFs, {len(real)} invoices, ${sum(i.total or 0 for i in real):,.2f} total, "
          f"{len(flagged)} to check -> {args.output}")


if __name__ == "__main__":
    main()
