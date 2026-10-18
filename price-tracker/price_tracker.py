"""Watch competitors' prices and stock, and write an Excel report of what changed.

    python price_tracker.py --pages 5 -o output/price_report.xlsx

Each run reads the listing pages of an online catalogue (here books.toscrape.com,
a practice shop built for scraping), saves every product's price and stock as a
dated snapshot CSV, compares it with the previous snapshot and writes an Excel
report: price drops and rises, new products, products gone and stock changes.
With --compare-only it skips the site and rebuilds the report from the newest
saved snapshot and the one before it.
"""

from __future__ import annotations

import argparse
import re
import time
from datetime import date, datetime
from pathlib import Path
from urllib.parse import urljoin

import pandas as pd
import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
SITE = "books.toscrape.com"
LISTING_URL = "https://books.toscrape.com/catalogue/page-{n}.html"
USER_AGENT = "PriceTrackerDemo/1.0 (portfolio demo; polite practice scraping of books.toscrape.com)"
TIMEOUT = 20     # seconds to wait for a page before giving up
MIN_DELAY = 1.0  # never ask for pages faster than one a second

RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}
SNAPSHOT_COLUMNS = ["url", "title", "price", "currency", "rating", "stock", "in_stock", "scraped_at"]
SNAPSHOT_NAME = "%Y-%m-%d_%H%M"
SNAPSHOT_FILE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}_\d{4}\.csv$")
NUMBER_RE = re.compile(r"\d[\d,]*(?:\.\d+)?")

# Friendly column names for product lists in the report.
PRODUCT_COLUMNS = {"title": "Product", "price": "Price", "rating": "Stars", "stock": "Stock", "url": "Link"}

MONEY = '"£"#,##0.00'
PERCENT = "0.0%"
COUNT = "#,##0"
GREY = "52514E"
HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
NOTE_FONT = Font(italic=True, color=GREY)
DROP_COLOUR = "2A78D6"
RISE_COLOUR = "E34948"
GRIDLINE_COLOUR = "D9D9D9"
TOP_MOVES = 10
LABEL_LENGTH = 30


# ---------- reading the listing pages ----------

def parse_price(text: str) -> tuple[float | None, str]:
    """'£51.77' -> (51.77, '£'). The currency is whatever comes before the number."""
    match = NUMBER_RE.search(text)
    if not match:
        return None, ""
    return float(match.group().replace(",", "")), text[:match.start()].strip()


def parse_rating(classes: list[str]) -> int | None:
    """Stars from the classes of p.star-rating: ['star-rating', 'Three'] -> 3."""
    return RATINGS.get(classes[1]) if len(classes) > 1 else None


def is_in_stock(stock: str) -> bool:
    return stock.strip().lower().startswith("in stock")


def _text(tag, selector: str) -> str:
    """Visible text of the first match, with runs of spaces and line breaks squeezed."""
    found = tag.select_one(selector)
    return " ".join(found.get_text(" ").split()) if found else ""


def parse_listing(html: str, page_url: str) -> tuple[list[dict], str | None]:
    """Products on one listing page, and the full URL of the next page (None on the last)."""
    soup = BeautifulSoup(html, "html.parser")
    products = []
    for pod in soup.select("article.product_pod"):
        link = pod.select_one("h3 a[title]")
        if link is None or not link.get("href"):
            continue
        price, currency = parse_price(_text(pod, "p.price_color"))
        stars = pod.select_one("p.star-rating")
        stock = _text(pod, "p.availability")
        products.append({
            "url": urljoin(page_url, link["href"]),
            "title": " ".join(link["title"].split()),  # the link text is cut short, the title isn't
            "price": price,
            "currency": currency,
            "rating": parse_rating(stars.get("class", [])) if stars else None,
            "stock": stock,
            "in_stock": is_in_stock(stock),
        })
    next_link = soup.select_one("li.next a[href]")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return products, next_url


def make_session() -> requests.Session:
    session = requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    return session


def fetch(session, url: str) -> str:
    response = session.get(url, timeout=TIMEOUT)
    response.raise_for_status()
    response.encoding = "utf-8"  # the site is UTF-8; without this "£" comes out as "Â£"
    return response.text


def scrape_catalogue(pages: int, delay: float, session=None, sleep=time.sleep) -> pd.DataFrame:
    """Read up to `pages` listing pages, waiting `delay` seconds between requests."""
    session = session or make_session()
    rows = []
    for n in range(1, pages + 1):
        if n > 1:
            sleep(delay)
        url = LISTING_URL.format(n=n)
        products, next_url = parse_listing(fetch(session, url), url)
        scraped_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        rows += [dict(product, scraped_at=scraped_at) for product in products]
        if next_url is None:
            break
    found = pd.DataFrame(rows, columns=SNAPSHOT_COLUMNS)
    found["rating"] = found["rating"].astype("Int64")
    return found.drop_duplicates("url").reset_index(drop=True)


# ---------- snapshots ----------

def snapshot_path(folder: Path, taken: datetime) -> Path:
    return folder / f"{taken:{SNAPSHOT_NAME}}.csv"


def snapshot_time(path: Path) -> datetime:
    return datetime.strptime(path.stem, SNAPSHOT_NAME)


def list_snapshots(folder: Path) -> list[Path]:
    """Snapshot files, oldest first. Names start with the date, so name order is time order."""
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.iterdir() if SNAPSHOT_FILE_RE.match(p.name))


def previous_snapshot(folder: Path, current: Path) -> Path | None:
    """The most recent snapshot saved before `current`, if there is one."""
    earlier = [p for p in list_snapshots(folder) if p.name < current.name]
    return earlier[-1] if earlier else None


def save_snapshot(products: pd.DataFrame, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    # utf-8-sig so "£" also shows correctly when the CSV is opened straight in Excel.
    products[SNAPSHOT_COLUMNS].to_csv(path, index=False, encoding="utf-8-sig", lineterminator="\n")
    return path


def load_snapshot(path: Path) -> pd.DataFrame:
    df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    missing = [c for c in SNAPSHOT_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(f"{path.name} is missing these columns: " + ", ".join(missing))
    df["price"] = pd.to_numeric(df["price"], errors="coerce")
    df["rating"] = pd.to_numeric(df["rating"], errors="coerce").astype("Int64")
    df["in_stock"] = df["in_stock"].str.strip().str.lower().isin(["true", "1", "yes"])
    return df[SNAPSHOT_COLUMNS].drop_duplicates("url").reset_index(drop=True)


# ---------- comparing two snapshots ----------

def product_table(df: pd.DataFrame, **labels) -> pd.DataFrame:
    """Snapshot rows with friendly column names; `labels` renames some, e.g. price="Last price"."""
    names = {column: labels.get(column, name) for column, name in PRODUCT_COLUMNS.items()}
    return df[list(names)].rename(columns=names).reset_index(drop=True)


def compare(old: pd.DataFrame, new: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """What changed between two snapshots. Products are matched on their URL."""
    both = new.merge(old, on="url", suffixes=("", "_old"))
    change = (both["price"] - both["price_old"]).round(2)
    price_changes = pd.DataFrame({
        "Product": both["title"],
        "Old price": both["price_old"],
        "New price": both["price"],
        "Change": change,
        "Change %": change / both["price_old"].where(both["price_old"] != 0),
        "Link": both["url"],
    })[change.abs() > 0]

    flipped = both[both["in_stock"] != both["in_stock_old"]]
    stock_changes = pd.DataFrame({
        "Product": flipped["title"],
        "Change": flipped["in_stock"].map({True: "Back in stock", False: "Out of stock"}),
        "Was": flipped["stock_old"],
        "Now": flipped["stock"],
        "Price": flipped["price"],
        "Link": flipped["url"],
    })

    new_products = new[~new["url"].isin(old["url"])].sort_values("title")
    gone = old[~old["url"].isin(new["url"])].sort_values("title")
    return {
        "price_changes": price_changes.sort_values(["Change %", "Product"]).reset_index(drop=True),
        "new_products": product_table(new_products),
        "gone": product_table(gone, price="Last price", stock="Last stock"),
        "stock_changes": stock_changes.sort_values(["Change", "Product"]).reset_index(drop=True),
    }


def summarise(current: pd.DataFrame, changes: dict | None) -> dict:
    summary = {"products": len(current), "average_price": current["price"].mean()}
    if changes is None:
        return summary
    change = changes["price_changes"]["Change"]
    stock = changes["stock_changes"]["Change"]
    summary.update({
        "drops": int((change < 0).sum()),
        "rises": int((change > 0).sum()),
        "new": len(changes["new_products"]),
        "gone": len(changes["gone"]),
        "back_in_stock": int((stock == "Back in stock").sum()),
        "out_of_stock": int((stock == "Out of stock").sum()),
    })
    return summary


def short_title(title: str, limit: int = LABEL_LENGTH) -> str:
    """Shorten long product names for chart labels, ending with '…'."""
    return title if len(title) <= limit else title[:limit - 1].rstrip() + "…"


def biggest_moves(price_changes: pd.DataFrame, count: int = TOP_MOVES) -> pd.DataFrame:
    """The largest % moves either way, biggest first, with names short enough for a chart."""
    size = price_changes["Change %"].abs().dropna().sort_values(ascending=False, kind="stable")
    top = price_changes.loc[size.index[:count]]
    return pd.DataFrame({
        "Product": top["Product"].map(short_title),
        "Old price": top["Old price"],
        "New price": top["New price"],
        "Change %": top["Change %"],
    }).reset_index(drop=True)


# ---------- writing Excel ----------

def _shown_length(value) -> int:
    """About how many characters a cell shows: prices and % have 2 decimals and a sign."""
    return len(f"{value:,.2f}") + 1 if isinstance(value, float) else len(str(value))


def write_table(ws, df: pd.DataFrame, start_row: int = 1, formats: dict | None = None) -> int:
    """Write df with a styled header at start_row; return the last row written."""
    formats = formats or {}
    for col, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, record in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col, value in enumerate(record, start=1):
            if pd.api.types.is_scalar(value) and pd.isna(value):
                value = None  # an empty cell, not "nan"
            cell = ws.cell(row=r, column=col, value=value)
            fmt = formats.get(df.columns[col - 1])
            if fmt:
                cell.number_format = fmt
    for col, name in enumerate(df.columns, start=1):
        width = max([len(str(name))] + [_shown_length(v) for v in df.iloc[:, col - 1].head(500)])
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 min(width + 3, 48))
    return start_row + len(df)


def colour_moves(ws, column: int, first_row: int, last_row: int) -> None:
    """Blue text for drops and red for rises, matching the chart."""
    for row in range(first_row, last_row + 1):
        cell = ws.cell(row=row, column=column)
        if isinstance(cell.value, (int, float)) and cell.value != 0:
            cell.font = Font(color=DROP_COLOUR if cell.value < 0 else RISE_COLOUR)


def link_cells(ws, column: int, first_row: int, last_row: int) -> None:
    """Make each product URL a clickable link."""
    for row in range(first_row, last_row + 1):
        cell = ws.cell(row=row, column=column)
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"


def write_sheet(wb, title: str, df: pd.DataFrame, empty_note: str, formats: dict, coloured=()):
    """A detail sheet: styled header, frozen top row, filters and clickable links."""
    ws = wb.create_sheet(title)
    last = write_table(ws, df, formats=formats)
    if df.empty:
        ws.cell(row=2, column=1, value=empty_note).font = NOTE_FONT
    for name in coloured:
        colour_moves(ws, df.columns.get_loc(name) + 1, 2, last)
    if "Link" in df.columns:
        link_cells(ws, df.columns.get_loc("Link") + 1, 2, last)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    return ws


def moves_chart(ws, top: pd.DataFrame, header_row: int) -> BarChart:
    """Horizontal bars of the biggest % moves: drops in blue, rises in red."""
    last_row = header_row + len(top)
    chart = BarChart()
    chart.type = "bar"  # horizontal bars leave room for product names
    chart.title = "Biggest price moves (% change)"
    chart.title.overlay = False  # keep the title above the bars
    chart.legend = None
    chart.height, chart.width = 9, 18
    chart.add_data(Reference(ws, min_col=4, min_row=header_row, max_row=last_row), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=header_row + 1, max_row=last_row))
    series = chart.series[0]
    series.invertIfNegative = False
    for idx, value in enumerate(top["Change %"]):
        colour = DROP_COLOUR if value < 0 else RISE_COLOUR
        point = DataPoint(idx=idx, invertIfNegative=False)
        point.graphicalProperties.solidFill = colour
        point.graphicalProperties.line.solidFill = colour
        series.dPt.append(point)
    chart.x_axis.scaling.orientation = "maxMin"  # biggest move at the top, as in the table
    chart.x_axis.tickLblPos = "low"              # names on the left, clear of bars below zero
    chart.y_axis.crosses = "max"                 # % scale along the bottom
    chart.y_axis.numFmt = "0%"
    chart.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRIDLINE_COLOUR))
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.gapWidth = 60
    return chart


def build_workbook(current, changes, summary, taken, compared_with=None, note=None) -> Workbook:
    """Summary first, then one sheet per kind of change, then every product."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    first_run = changes is None
    tables = compare(current.head(0), current.head(0)) if first_run else changes  # no rows, same columns
    empty = "No earlier snapshot to compare with yet." if first_run else "None since the last snapshot."
    write_sheet(wb, "Price changes", tables["price_changes"], empty,
                {"Old price": MONEY, "New price": MONEY, "Change": MONEY, "Change %": PERCENT},
                coloured=("Change", "Change %"))
    write_sheet(wb, "New products", tables["new_products"], empty, {"Price": MONEY})
    write_sheet(wb, "Gone", tables["gone"], empty, {"Last price": MONEY})
    write_sheet(wb, "Stock changes", tables["stock_changes"], empty, {"Price": MONEY})
    write_sheet(wb, "All products", product_table(current), "No products in this snapshot.", {"Price": MONEY})

    ws["A1"] = f"Competitor price report: {SITE}"
    ws["A1"].font = Font(bold=True, size=16)
    if first_run:
        ws["A2"] = (f"First snapshot ({taken:%d %b %Y %H:%M}), so there is nothing earlier to compare "
                    f"with yet. Built {date.today():%d %b %Y}.")
    else:
        ws["A2"] = (f"Snapshot of {taken:%d %b %Y %H:%M} compared with {compared_with:%d %b %Y %H:%M}. "
                    f"Built {date.today():%d %b %Y}.")
    ws["A2"].font = Font(color=GREY)
    if note:
        ws["A3"] = note
        ws["A3"].font = NOTE_FONT

    kpis = [
        ("Products tracked", summary["products"], COUNT),
        ("Average price", summary["average_price"], MONEY),
    ]
    if not first_run:
        kpis += [
            ("Price drops", summary["drops"], COUNT),
            ("Price rises", summary["rises"], COUNT),
            ("New products", summary["new"], COUNT),
            ("Products gone", summary["gone"], COUNT),
            ("Back in stock", summary["back_in_stock"], COUNT),
            ("Went out of stock", summary["out_of_stock"], COUNT),
        ]
    row = 5 if note else 4
    for offset, (label, value, fmt) in enumerate(kpis):
        ws.cell(row=row + offset, column=1, value=label)
        cell = ws.cell(row=row + offset, column=2, value=None if pd.isna(value) else value)
        cell.number_format = fmt
        cell.font = Font(bold=True)

    row += len(kpis) + 1
    ws.cell(row=row, column=1, value="Biggest moves").font = Font(bold=True, size=12)
    if tables["price_changes"].empty:
        message = empty if first_run else "No price changes since the last snapshot."
        ws.cell(row=row + 1, column=1, value=message).font = NOTE_FONT
    else:
        top = biggest_moves(tables["price_changes"])
        last = write_table(ws, top, start_row=row + 1,
                           formats={"Old price": MONEY, "New price": MONEY, "Change %": PERCENT})
        colour_moves(ws, 4, row + 2, last)
        ws.add_chart(moves_chart(ws, top, header_row=row + 1), "F4")
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 14
    return wb


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--pages", type=int, default=5,
                        help="listing pages to read, 20 products each (default 5)")
    parser.add_argument("--delay", type=float, default=1.0,
                        help="seconds to wait between page requests, at least 1 (default 1.0)")
    parser.add_argument("--snapshots", type=Path, default=HERE / "snapshots",
                        help="folder for the dated snapshot CSVs (default: snapshots/ next to this script)")
    parser.add_argument("-o", "--output", type=Path, default=HERE / "output" / "price_report.xlsx",
                        help="Excel report to write (default: output/price_report.xlsx next to this script)")
    parser.add_argument("--compare-only", action="store_true",
                        help="don't visit the site: compare the newest saved snapshot with the one before")
    parser.add_argument("--note", help="an extra line to show under the report title")
    args = parser.parse_args(argv)
    if args.pages < 1:
        parser.error("--pages must be at least 1")
    if args.delay < MIN_DELAY:
        parser.error(f"--delay must be at least {MIN_DELAY:g} second, to be polite to the site")

    if args.compare_only:
        saved = list_snapshots(args.snapshots)
        if not saved:
            raise SystemExit(f"No snapshots in {args.snapshots} yet. Run once without --compare-only first.")
        current_path = saved[-1]
    else:
        taken = datetime.now()
        try:
            products = scrape_catalogue(args.pages, args.delay)
        except requests.RequestException as error:
            raise SystemExit(f"Could not read the catalogue ({error}). Nothing was saved.")
        if products.empty:
            # Saving an empty snapshot would make every product look "gone" next time.
            raise SystemExit("No products found on the listing pages (has the site changed?). Nothing was saved.")
        current_path = save_snapshot(products, snapshot_path(args.snapshots, taken))

    current = load_snapshot(current_path)
    previous_path = previous_snapshot(args.snapshots, current_path)
    changes = compare(load_snapshot(previous_path), current) if previous_path else None
    summary = summarise(current, changes)
    wb = build_workbook(current, changes, summary, snapshot_time(current_path),
                        snapshot_time(previous_path) if previous_path else None, args.note)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(args.output)
    except PermissionError:
        raise SystemExit(f"Could not write {args.output}. Is it open in Excel? "
                         "Close it and run again with --compare-only.")

    compared = f"compared with {previous_path.stem}" if previous_path else "nothing earlier to compare with yet"
    print(f"{summary['products']:,} products tracked, average price £{summary['average_price']:,.2f} "
          f"(snapshot {current_path.stem}, {compared}).")
    if changes is None:
        print(f"Changes will show from the next run. Report: {args.output}")
    else:
        s = summary
        print(f"{s['drops'] + s['rises']} price changes ({s['drops']} down, {s['rises']} up), "
              f"{s['new']} new, {s['gone']} gone, {s['back_in_stock']} back in stock, "
              f"{s['out_of_stock']} out of stock. Report: {args.output}")


if __name__ == "__main__":
    main()
