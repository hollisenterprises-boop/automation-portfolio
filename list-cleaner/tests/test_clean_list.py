import sys
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import clean_list as cl  # noqa: E402
import make_sample_data as samples  # noqa: E402


@pytest.mark.parametrize("raw, want", [
    ("Olivia Smith", ("Olivia", "Smith")),
    ("  OLIVIA   SMITH ", ("Olivia", "Smith")),
    ("smith, olivia", ("Olivia", "Smith")),
    ("Mrs Olivia Smith", ("Olivia", "Smith")),
    ("Dr. james o'brien", ("James", "O'Brien")),
    ("ruby mcdonald", ("Ruby", "McDonald")),
    ("Ella Smith-Jones", ("Ella", "Smith-Jones")),
    ("DeShawn McAllister", ("DeShawn", "McAllister")),  # mixed case is trusted
    ("", ("", "")),
])
def test_clean_name(raw, want):
    assert cl.clean_name(raw) == want


@pytest.mark.parametrize("raw, email, problem", [
    ("Olivia.Smith@Example.COM", "olivia.smith@example.com", False),
    (" mailto:olivia@example.org ", "olivia@example.org", False),
    ("olivia@example", "", True),
    ("olivia smith@example.com", "", True),
    ("see notes", "", True),
    ("", "", False),
])
def test_clean_email(raw, email, problem):
    got, issue = cl.clean_email(raw)
    assert got == email
    assert (issue is not None) == problem


@pytest.mark.parametrize("raw, want", [
    ("(08) 5550 1234", "08 5550 1234"),
    ("+61 8 5550 1234", "08 5550 1234"),
    ("0855501234", "08 5550 1234"),
    ("08 5550-1234", "08 5550 1234"),
    ("855501234", "08 5550 1234"),  # leading zero dropped by Excel
    ("0491 570 156", "0491 570 156"),
    ("+61 491 570 156", "0491 570 156"),
    ("61491570156", "0491 570 156"),
])
def test_clean_phone(raw, want):
    assert cl.clean_phone(raw) == (want, None)


@pytest.mark.parametrize("raw", ["0855 501", "12345", "+1 555 0100 2000"])
def test_bad_phone_is_flagged(raw):
    phone, problem = cl.clean_phone(raw)
    assert phone == "" and "not a valid Australian number" in problem


@pytest.mark.parametrize("suburb, state, want", [
    ("Perth, WA 6000", "", ("Perth", "WA", "6000")),
    ("PERTH", "Western Australia", ("Perth", "WA", "")),
    ("gold coast", "qld", ("Gold Coast", "QLD", "")),
    ("Fremantle WA", "", ("Fremantle", "WA", "")),
])
def test_clean_place(suburb, state, want):
    assert cl.clean_place(suburb, state) == want


def rec(**kw):
    defaults = dict(source="a.csv", row=2, raw={})
    defaults.update(kw)
    return cl.Record(**defaults)


def test_one_no_means_no():
    members = [rec(first="Ava", last="Lee", consent=True), rec(first="Ava", last="Lee", consent=False)]
    assert cl.merge(members, cl.EXACT).marketing == "No"
    assert cl.merge(members[:1], cl.EXACT).marketing == "Yes"
    assert cl.merge([rec(first="Ava", last="Lee")], cl.EXACT).marketing == "Unknown"


def test_keeps_the_full_first_name():
    members = [rec(first="Kate", last="Lee"), rec(first="Katherine", last="Lee")]
    assert cl.merge(members, cl.SIMILAR).first == "Katherine"


@pytest.mark.parametrize("a, b, same", [
    ("Kate", "Katherine", True), ("Sam", "Samantha", True), ("Sam", "Samuel", True),
    ("Samuel", "Samantha", False), ("Jon", "John", True), ("Mia", "Noah", False),
])
def test_same_first_name(a, b, same):
    assert cl.same_first_name(a, b) is same


def test_merges_on_email_or_phone():
    records = [rec(first="Ava", last="Lee", email="ava@example.com"),
               rec(first="A", last="Lee", email="ava@example.com", phone="08 5550 1111"),
               rec(first="Ava", last="Lee", phone="08 5550 1111"),
               rec(first="Ava", last="Lee", email="other@example.com")]
    groups = cl.group_records(records)
    assert sorted(len(m) for m, _ in groups) == [1, 3]


def test_similar_names_with_different_phones_stay_apart():
    records = [rec(first="Sam", last="Lee", suburb="Perth", phone="08 5550 1111"),
               rec(first="Samantha", last="Lee", suburb="Perth", phone="08 5550 2222"),
               rec(first="Samantha", last="Lee", suburb="Perth")]
    sizes = sorted(len(m) for m, _ in cl.group_records(records))
    assert sizes == [1, 2]  # the record with no phone joins one of them, never both


@pytest.fixture(scope="module")
def sample(tmp_path_factory):
    folder = tmp_path_factory.mktemp("lists")
    samples.main(folder)
    return cl.clean(folder)


def test_sample_finds_every_real_person(sample):
    contacts, records, rows_per_file, blanks = sample
    assert sum(rows_per_file.values()) == 197
    assert blanks == 2
    assert len(contacts) == 138  # the number of real people the sample was built from


def test_sample_merge_keeps_the_good_email(sample):
    contacts = sample[0]
    james = [c for c in contacts if (c.first, c.last, c.suburb) == ("James", "Wilson", "Brisbane")]
    assert len(james) == 1
    assert james[0].email == "james.wilson74@example.org"  # the shop export's "@example" was broken
    assert james[0].phone == "07 5550 1596"


def test_sample_problems_listed(sample):
    problems = [p for r in sample[1] for p in r.problems]
    assert len(problems) == 4


def test_workbook(sample, tmp_path):
    path = tmp_path / "out.xlsx"
    cl.build_workbook(*sample).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Clean List", "Merged Rows", "To Check"]
    assert wb["Clean List"].max_row == 1 + 138
    assert wb["Summary"]["B4"].value == 138
    assert wb["To Check"].max_row == 1 + 4
