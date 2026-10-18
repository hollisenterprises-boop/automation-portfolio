import io
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import sales_report as sr  # noqa: E402

SAMPLE = """\
 Order ID,Order Date ,SKU,Product Name,Qty,Item Price,Order Status,Ship State
SS-1,2026-04-01,mp-set-5 ,Glass Meal Prep Set (5 pack),1,$39.95,shipped,Western Australia
SS-1,01/04/2026,MP-SET-5,Glass Meal Prep Set (5 pack),1 unit,AUD 39.95,Shipped,WA
SS-2,2026-04-02,LUNCH-2C,Two-Compartment Lunch Box,2,24.95,CANCELED,nsw

SS-3,3 Apr 2026,LUNCH-2C,Two-Compartment Lunch Box,1,TBC,Delivered,VIC
SS-4,2026-05-10,LUNCH-2C,Two-Compartment Lunch Box,x2,24.95,Refunded,QLD
SS-5,2026-05-11 09:15:00,GLS-JAR-500,Glass  Storage Jar 500ml,3,14.95,Delivered,w.a.
Total rows: 7,,,,,,,
"""


def cleaned():
    return sr.clean_orders(sr.load_export(io.StringIO(SAMPLE)))


@pytest.mark.parametrize("text", ["2026-04-03", "03/04/2026", "3 Apr 2026",
                                  "2026-04-03 14:22:10", "April 3 2026", " 2026-04-03 "])
def test_dates_in_every_format(text):
    assert sr.parse_date(text) == date(2026, 4, 3)


def test_unreadable_dates():
    assert sr.parse_date("") is None
    assert sr.parse_date("soon") is None


@pytest.mark.parametrize("text", ["24.95", "$24.95", "AUD 24.95", " 24.95 ", "A$24.95"])
def test_prices_with_currency_text(text):
    assert sr.parse_money(text) == 24.95


def test_price_edge_cases():
    assert sr.parse_money("$1,249.00") == 1249.0
    assert sr.parse_money("TBC") is None
    assert sr.parse_money("") is None


def test_quantities():
    assert [sr.parse_qty(t) for t in ["2", "2 units", "x3", "1 unit"]] == [2, 2, 3, 1]
    assert sr.parse_qty("") is None


def test_status_and_state_names():
    assert sr.normalize_status("SHIPPED") == "Shipped"
    assert sr.normalize_status("delivered") == "Delivered"
    assert sr.normalize_status("canceled") == "Cancelled"
    assert sr.normalize_status("pending ") == "Pending"
    assert sr.normalize_status("lost") == "Unknown"
    assert sr.normalize_state("W.A.") == "WA"
    assert sr.normalize_state("new south wales") == "NSW"
    assert sr.normalize_state("Mars") == "Unknown"


def test_cleaning_removes_junk_and_keeps_real_orders():
    clean, fixes = cleaned()
    assert list(clean["Order ID"]) == ["SS-1", "SS-2", "SS-4", "SS-5"]
    removed = dict(zip(fixes["Row"], fixes["What happened"]))
    assert removed[3] == "Exact duplicate of an earlier row, removed"
    assert removed[5] == "Blank row removed"
    assert removed[6] == "Could not read this value, row left out"
    assert removed[9] == "Not an order (summary or footer line), removed"
    jar = clean[clean["Order ID"] == "SS-5"].iloc[0]
    assert jar["Product Name"] == "Glass Storage Jar 500ml"
    assert jar["Ship State"] == "WA"
    assert jar["SKU"] == "GLS-JAR-500"
    # Correct values are left alone and not reported as fixes.
    assert not ((fixes["Row"] == 8) & (fixes["Column"] == "Order Status")).any()
    assert set(fixes.loc[fixes["What happened"].isin(sr.REMOVALS), "Row"]) == {3, 5, 6, 9}


def test_missing_column_is_explained():
    with pytest.raises(ValueError, match="Ship State"):
        sr.clean_orders(sr.load_export(io.StringIO("Order ID,Order Date,SKU,Product Name,"
                                                   "Qty,Item Price,Order Status\n")))


def test_summary_counts_only_shipped_sales():
    clean, _ = cleaned()
    summary = sr.summarise(clean)
    assert summary["sales"] == pytest.approx(39.95 + 3 * 14.95)
    assert summary["orders"] == 2
    assert summary["units"] == 4
    assert summary["refunded"] == pytest.approx(49.90)
    assert summary["cancelled_orders"] == 1
    assert list(summary["by_month"]["Month"]) == ["Apr 2026", "May 2026"]


def test_workbook_has_every_sheet(tmp_path):
    clean, fixes = cleaned()
    path = tmp_path / "report.xlsx"
    sr.build_workbook(clean, fixes, sr.summarise(clean), "Test Store", "test.csv").save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "By Month", "By Product", "By State",
                             "Clean Data", "Fixes Log"]
    assert wb["Summary"]["B4"].value == pytest.approx(84.80)
