"""Load a store's orders into SQLite, run SQL reports on them, and write one formatted Excel workbook.

    python sql_reports.py -o output --as-of 2026-10-19

Reads the order export from the sales report demo (cleaned by that demo's code)
and a customer list, builds a SQLite database from schema.sql, then runs every
.sql file in queries/ and writes the results to Excel: one sheet for each
query, a summary page with charts, and a sheet showing the SQL behind each
report. To add a report, add a .sql file with a "-- title:" line.
"""

from __future__ import annotations

import argparse
import csv
import re
import sqlite3
import sys
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent / "sales-report"))
import sales_report  # noqa: E402  (its cleaning step reads the messy order export)

SCHEMA = HERE / "schema.sql"
QUERIES = HERE / "queries"
ORDERS = HERE.parent / "sales-report" / "sample_data" / "orders_export_messy.csv"
CUSTOMERS = HERE / "sample_data" / "customers.csv"
ORDER_CUSTOMERS = HERE / "sample_data" / "order_customers.csv"
HEADLINE = "00_headline.sql"  # its one row becomes the Summary sheet's numbers

MONEY = '"$"#,##0.00'
COUNT = "#,##0"
PERCENT = "0.0%"
CHANGE = "+0.0%;-0.0%;0.0%"
COLUMN_FORMATS = {
    "Orders": COUNT, "Units": COUNT, "Units sold": COUNT, "Customers": COUNT, "Customers who bought": COUNT,
    "Repeat customers": COUNT, "Days between first and latest": COUNT,
    "Sales": MONEY, "Net sales": MONEY, "Average order": MONEY, "Total spent": MONEY,
    "Share of sales": PERCENT, "Running share": PERCENT, "Repeat customer rate": PERCENT,
    "Share of sales from repeat customers": PERCENT, "Change on previous month": CHANGE,
}
DATE_COLUMNS = {"First order", "Latest order"}  # the database keeps them as YYYY-MM-DD text
BAR_COLUMNS = {"Sales", "Total spent"}  # get an in-cell bar
HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
BAR_COLOUR = "2A78D6"


@dataclass
class Report:
    file: str
    title: str
    about: str
    sql: str
    columns: list[str]
    rows: list[tuple]


# ---------- building the database ----------

def read_csv(path: Path, columns: list[str]) -> list[tuple]:
    with Path(path).open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        missing = [c for c in columns if c not in (reader.fieldnames or [])]
        if missing:
            raise ValueError(f"{Path(path).name} is missing these columns: {', '.join(missing)}")
        return [tuple(record[c] for c in columns) for record in reader]


def create_schema(con: sqlite3.Connection) -> None:
    con.execute("PRAGMA foreign_keys = ON")
    con.executescript(SCHEMA.read_text(encoding="utf-8"))


def build_database(db_path, orders_csv=ORDERS, customers_csv=CUSTOMERS,
                   order_customers_csv=ORDER_CUSTOMERS) -> sqlite3.Connection:
    """Clean the order export, load it into a new SQLite database and return the open connection."""
    clean, _fixes = sales_report.clean_orders(sales_report.load_export(orders_csv))
    customers = read_csv(customers_csv, ["Customer ID", "Name", "State"])
    owner = dict(read_csv(order_customers_csv, ["Order ID", "Customer ID"]))

    order_ids = list(dict.fromkeys(clean["Order ID"]))
    missing = [o for o in order_ids if o not in owner]
    if missing:
        raise ValueError(f"{len(missing)} orders have no customer in {Path(order_customers_csv).name}, "
                         f"for example {missing[0]}.")
    unknown = sorted({owner[o] for o in order_ids} - {c[0] for c in customers})
    if unknown:
        raise ValueError(f"{Path(customers_csv).name} has no customer {unknown[0]} "
                         f"({len(unknown)} customer IDs are missing).")
    for column in ("Order Date", "Order Status", "Ship State"):
        if clean.groupby("Order ID")[column].nunique().max() > 1:
            raise ValueError(f"Some orders have lines with different values in {column}.")

    if str(db_path) != ":memory:":
        Path(db_path).parent.mkdir(parents=True, exist_ok=True)
        Path(db_path).unlink(missing_ok=True)
    con = sqlite3.connect(db_path)
    create_schema(con)

    names = clean.groupby("SKU")["Product Name"].agg(lambda names: names.mode().iat[0])
    first = clean.drop_duplicates("Order ID")
    con.executemany("INSERT INTO customers VALUES (?, ?, ?)", customers)
    con.executemany("INSERT INTO products VALUES (?, ?)", names.items())
    con.executemany("INSERT INTO orders VALUES (?, ?, ?, ?, ?)", [
        (order_id, day, status, state, owner[order_id]) for order_id, day, status, state in zip(
            first["Order ID"], first["Order Date"].dt.strftime("%Y-%m-%d"),
            first["Order Status"], first["Ship State"])])
    con.executemany("INSERT INTO order_items VALUES (?, ?, ?, ?)", [
        (order_id, sku, int(qty), round(price * 100)) for order_id, sku, qty, price in zip(
            clean["Order ID"], clean["SKU"], clean["Qty"], clean["Item Price"])])
    con.commit()
    return con


# ---------- running the queries ----------

def header_value(sql: str, key: str) -> str | None:
    match = re.search(rf"^--\s*{key}:\s*(.+?)\s*$", sql, re.I | re.M)
    return match.group(1) if match else None


def run_report(con: sqlite3.Connection, path: Path) -> Report:
    """Run one .sql file. Its "-- title:" and "-- about:" comment lines name the sheet and say what it shows."""
    sql = Path(path).read_text(encoding="utf-8")
    cursor = con.execute(sql)
    title = header_value(sql, "title") or Path(path).stem
    if len(title) > 31 or re.search(r"[\[\]:*?/\\]", title):
        raise ValueError(f"{Path(path).name}: the title \"{title}\" can't be an Excel sheet name "
                         "(31 characters at most, none of []:*?/\\).")
    return Report(Path(path).name, title, header_value(sql, "about") or "", sql.strip(),
                  [d[0] for d in cursor.description], cursor.fetchall())


def run_all(con: sqlite3.Connection, queries_dir: Path = QUERIES) -> list[Report]:
    return [run_report(con, path) for path in sorted(Path(queries_dir).glob("[0-9][0-9]_*.sql"))]


# ---------- writing Excel ----------

def present(column: str, value):
    if value is None:
        return None
    if column == "Month":
        return datetime.strptime(value, "%Y-%m").strftime("%b %Y")
    if column in DATE_COLUMNS:
        return date.fromisoformat(value)
    return value


def write_table(ws, columns: list[str], rows: list[tuple]) -> None:
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    shown = [[present(column, value) for column, value in zip(columns, row)] for row in rows]
    for r, record in enumerate(shown, start=2):
        for col, value in enumerate(record, start=1):
            cell = ws.cell(row=r, column=col, value=value)
            if columns[col - 1] in COLUMN_FORMATS:
                cell.number_format = COLUMN_FORMATS[columns[col - 1]]
            elif isinstance(value, date):
                cell.number_format = "dd/mm/yyyy"
    for col, name in enumerate(columns, start=1):
        width = max([len(name) * 0.8] + [len(str(record[col - 1])) for record in shown[:500]])
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 3, 11), 48)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, name in enumerate(columns, start=1):
        if name in BAR_COLUMNS and rows:
            letter = get_column_letter(col)
            ws.conditional_formatting.add(f"{letter}2:{letter}{len(rows) + 1}", DataBarRule(
                start_type="num", start_value=0, end_type="max", color=BAR_COLOUR))


def add_bar_chart(summary, sheet, report: Report, category: str, value: str, title: str, anchor: str,
                  horizontal: bool = False) -> None:
    n = len(report.rows)
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.title = title
    chart.legend = None
    chart.height, chart.width = 8, 16
    chart.add_data(Reference(sheet, min_col=report.columns.index(value) + 1, min_row=1, max_row=n + 1),
                   titles_from_data=True)
    chart.set_categories(Reference(sheet, min_col=report.columns.index(category) + 1, min_row=2, max_row=n + 1))
    chart.series[0].graphicalProperties.solidFill = BAR_COLOUR
    chart.series[0].graphicalProperties.line.solidFill = BAR_COLOUR
    chart.y_axis.numFmt = '"$"#,##0'
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.gapWidth = 60
    if horizontal:
        chart.x_axis.scaling.orientation = "maxMin"  # biggest at the top
        chart.y_axis.crosses = "max"  # and keep the dollar scale along the bottom
    summary.add_chart(chart, anchor)


def build_workbook(reports: list[Report], store: str, source_name: str, as_of: date) -> Workbook:
    headline = next(r for r in reports if r.file == HEADLINE)
    numbers = dict(zip(headline.columns, headline.rows[0]))
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    sheets = {}
    for report in reports:
        if report.file == HEADLINE:
            continue
        ws = wb.create_sheet(report.title)
        write_table(ws, report.columns, report.rows)
        sheets[report.title] = (ws, report)

    queries = wb.create_sheet("Queries")
    for col, name in enumerate(["Sheet", "File", "What it answers", "SQL"], start=1):
        cell = queries.cell(row=1, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for r, report in enumerate(reports, start=2):
        queries.cell(row=r, column=1, value=report.title)
        queries.cell(row=r, column=2, value=report.file)
        queries.cell(row=r, column=3, value=report.about)
        sql = queries.cell(row=r, column=4, value=report.sql)
        sql.font = Font(name="Consolas", size=10)
        for col in (1, 2, 3, 4):
            queries.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
        queries.row_dimensions[r].height = min(409, 14 * (report.sql.count("\n") + 1) + 6)
    for letter, width in zip("ABCD", (22, 26, 44, 110)):
        queries.column_dimensions[letter].width = width
    queries.freeze_panes = "A2"

    span = ""
    if numbers["First order"]:
        first, last = date.fromisoformat(numbers["First order"]), date.fromisoformat(numbers["Latest order"])
        span = f"{first:%d %b %Y} to {last:%d %b %Y}. "
    summary["A1"] = f"SQL sales report: {store}"
    summary["A1"].font = Font(bold=True, size=16)
    summary["A2"] = (f"{span}{len(reports)} SQL queries on a SQLite database "
                     f"built from {source_name} on {as_of:%d %b %Y}.")
    summary["A2"].font = Font(color="52514E")
    kpis = [c for c in headline.columns if c not in DATE_COLUMNS]
    for offset, label in enumerate(kpis):
        summary.cell(row=4 + offset, column=1, value=label)
        cell = summary.cell(row=4 + offset, column=2, value=numbers[label])
        cell.number_format = COLUMN_FORMATS.get(label, "General")
        cell.font = Font(bold=True)
    summary.column_dimensions["A"].width = 38
    summary.column_dimensions["B"].width = 16
    summary["A13"] = ("Sales are shipped, delivered and completed orders; cancelled, refunded and pending "
                      "orders are left out. All data, including the customers, is made up.")
    summary.merge_cells("A13:B17")
    summary["A13"].font = Font(color="52514E", italic=True)
    summary["A13"].alignment = Alignment(wrap_text=True, vertical="top")

    if "Sales by month" in sheets:
        ws, report = sheets["Sales by month"]
        add_bar_chart(summary, ws, report, "Month", "Sales", "Net sales by month", "D4")
    if "Totals by state" in sheets:
        ws, report = sheets["Totals by state"]
        add_bar_chart(summary, ws, report, "State", "Sales", "Net sales by state", "D21", horizontal=True)
    return wb


def main(argv=None) -> None:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("-o", "--output", type=Path, default=Path("output"), help="folder for the database and the report")
    parser.add_argument("--orders", type=Path, default=ORDERS, help="order export CSV (default: the sales report demo's)")
    parser.add_argument("--customers", type=Path, default=CUSTOMERS, help="customers CSV")
    parser.add_argument("--order-customers", type=Path, default=ORDER_CUSTOMERS, help="CSV saying which customer placed each order")
    parser.add_argument("--store", default="Sample Store", help="name shown on the report")
    parser.add_argument("--as-of", type=date.fromisoformat, default=None, metavar="YYYY-MM-DD",
                        help="date shown on the report (default: today)")
    args = parser.parse_args(argv)

    args.output.mkdir(parents=True, exist_ok=True)
    con = build_database(args.output / "store.db", args.orders, args.customers, args.order_customers)
    reports = run_all(con)
    con.close()
    path = args.output / "sql_report.xlsx"
    build_workbook(reports, args.store, args.orders.name, args.as_of or date.today()).save(path)

    headline = next(r for r in reports if r.file == HEADLINE)
    numbers = dict(zip(headline.columns, headline.rows[0]))
    print(f"Net sales ${numbers['Net sales']:,.2f} from {numbers['Orders']:,} orders; "
          f"{numbers['Repeat customers']:,} of {numbers['Customers who bought']:,} customers ordered again.")
    print(f"{len(reports)} queries run. Database: {args.output / 'store.db'}  Report: {path}")


if __name__ == "__main__":
    main()
