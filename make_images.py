"""Draw one 1600x1200 portfolio image per demo from its committed sample output.

    python make_images.py            # all demos
    python make_images.py sales      # just one

Images go to images/. Each shows the messy input next to the finished result.
"""

import csv
import re
import sqlite3
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.patches import FancyBboxPatch  # noqa: E402
from matplotlib.ticker import FuncFormatter  # noqa: E402

HERE = Path(__file__).parent
OUT = HERE / "images"

SURFACE = "#fcfcfb"
INK = "#0b0b0b"
INK_2 = "#52514e"
MUTED = "#898781"
GRID = "#e1e0d9"
BASELINE = "#c3c2b7"
TILE = "#f3f2ee"
BLUE = "#2a78d6"
RED = "#e34948"
FIXED_WASH = "#fde3d6"
STATUS = {"good": "#0ca30c", "warning": "#fab219", "serious": "#ec835a", "critical": "#d03b3b"}

plt.rcParams.update({
    "font.family": ["Segoe UI", "DejaVu Sans"],
    "text.color": INK,
    "axes.edgecolor": BASELINE,
    "axes.labelcolor": INK_2,
    "xtick.color": MUTED,
    "ytick.color": MUTED,
})


def canvas(title, subtitle):
    fig = plt.figure(figsize=(16, 12), dpi=100, facecolor=SURFACE)
    fig.text(0.04, 0.945, title, fontsize=34, fontweight="bold", va="top")
    fig.text(0.04, 0.885, subtitle, fontsize=18, color=INK_2, va="top")
    return fig


def footer(fig, tools, note):
    fig.text(0.04, 0.03, tools, fontsize=13, color=MUTED)
    fig.text(0.96, 0.03, note, fontsize=13, color=MUTED, ha="right")


def section_label(fig, x, y, text):
    fig.text(x, y, text, fontsize=16, fontweight="bold", color=INK_2, va="bottom")


def kpi_tiles(fig, left, bottom, width, height, tiles):
    """tiles: list of (value text, label text)."""
    gap = 0.015
    tile_w = (width - gap * (len(tiles) - 1)) / len(tiles)
    for i, (value, label) in enumerate(tiles):
        ax = fig.add_axes([left + i * (tile_w + gap), bottom, tile_w, height])
        ax.axis("off")
        ax.add_patch(FancyBboxPatch((0, 0), 1, 1, boxstyle="round,pad=0,rounding_size=0.06",
                                    transform=ax.transAxes, facecolor=TILE, edgecolor="none"))
        ax.text(0.08, 0.6, value, fontsize=26, fontweight="bold", transform=ax.transAxes, va="center")
        ax.text(0.08, 0.24, label, fontsize=13, color=INK_2, transform=ax.transAxes, va="center")


def raw_table(fig, rect, header, rows, highlight, col_widths, font_size=13):
    """Draw raw export rows; highlight is a set of (row index, col index), col -1 = whole row."""
    ax = fig.add_axes(rect)
    ax.axis("off")
    table = ax.table(cellText=rows, colLabels=header, cellLoc="left", loc="upper left",
                     colWidths=col_widths, bbox=[0, 0, 1, 1])
    table.auto_set_font_size(False)
    for (r, c), cell in table.get_celld().items():
        cell.PAD = 0.05 if col_widths[c] > 0.3 else 0.1  # same text inset in wide and narrow columns
        cell.set_edgecolor(GRID)
        cell.set_linewidth(0.8)
        if r == 0:
            cell.set_facecolor("#184f95")
            cell.get_text().set_color("white")
            cell.get_text().set_fontweight("bold")
            cell.get_text().set_fontsize(font_size)
            continue
        cell.get_text().set_fontfamily(["Consolas", "DejaVu Sans Mono"])
        cell.get_text().set_fontsize(font_size)
        if (r - 1, c) in highlight or (r - 1, -1) in highlight:
            cell.set_facecolor(FIXED_WASH)
        else:
            cell.set_facecolor("white")
    return ax


def style_axis(ax):
    ax.set_facecolor(SURFACE)
    for side in ("top", "right", "left"):
        ax.spines[side].set_visible(False)
    ax.spines["bottom"].set_color(BASELINE)
    ax.tick_params(length=0, labelsize=13)
    ax.grid(axis="y", color=GRID, linewidth=1)
    ax.set_axisbelow(True)


# ---------- sales report ----------

FIX_LABELS = [
    ("Date written as YYYY-MM-DD", "dates in mixed formats"),
    ("Currency text removed", "prices with $ or AUD stuck on"),
    ("State written as its code", "state names written out"),
    ("Status written one way", "status spellings (shipped, SHIPPED...)"),
    ("Spaces and capitals fixed", "SKUs with stray spaces or lower case"),
    ("Quantity read as a number", "quantities like \"3 units\""),
    ("Extra spaces removed", "product names with double spaces"),
    ("Exact duplicate of an earlier row, removed", "duplicate rows removed"),
    ("Blank row removed", "blank rows removed"),
    ("Could not read this value, row left out", "unreadable rows set aside (price \"TBC\")"),
    ("Not an order (summary or footer line), removed", "footer line removed"),
]
WHOLE_ROW_REMOVALS = {"Exact duplicate of an earlier row, removed",
                      "Not an order (summary or footer line), removed"}

def sales_image():
    demo = HERE / "sales-report"
    report = demo / "output" / "sales_report.xlsx"
    with (demo / "sample_data" / "orders_export_messy.csv").open(encoding="utf-8") as f:
        lines = list(csv.reader(f))
    header = [h.strip() for h in lines[0]]
    keep = ["Order Date", "SKU", "Qty", "Item Price", "Order Status", "Ship State"]
    idx = [header.index(k) for k in keep]
    sample = lines[1:12]  # file lines 2-12
    rows = [[r[i] if r else "" for i in idx] for r in sample]
    for n, r in enumerate(sample):
        if not r:
            rows[n][0] = "(blank row)"

    fixes = pd.read_excel(report, sheet_name="Fixes Log")
    highlight = set()
    for _, fix in fixes[fixes["Row"].between(2, 12)].iterrows():
        row = int(fix["Row"]) - 2
        col = fix["Column"] if isinstance(fix["Column"], str) else ""
        if col in keep:
            highlight.add((row, keep.index(col)))
        elif fix["What happened"] in WHOLE_ROW_REMOVALS:
            highlight.add((row, -1))
    for n, r in enumerate(sample):
        if not r:
            highlight.add((n, -1))

    summary = load_summary(report)
    months = pd.read_excel(report, sheet_name="By Month")

    fig = canvas("Sales report builder", "A messy store export in, a clean Excel report out, in seconds")

    section_label(fig, 0.04, 0.80, "Before: the store's raw export")
    raw_table(fig, [0.04, 0.40, 0.50, 0.39], keep, rows, highlight,
              [0.265, 0.165, 0.095, 0.13, 0.135, 0.21], font_size=12)
    fig.text(0.04, 0.365, "Shaded = fixed or removed. Every change is listed in the Fixes Log.",
             fontsize=13, color=INK_2)

    counts = fixes["What happened"].value_counts()
    found = [(counts.get(key, 0), label) for key, label in FIX_LABELS]
    fig.text(0.04, 0.315, "What it fixed in this export", fontsize=15, fontweight="bold", color=INK_2)
    for n, (count, label) in enumerate(c for c in found if c[0]):
        y = 0.28 - n * 0.019
        fig.text(0.085, y, f"{count:,}", fontsize=13, fontweight="bold", ha="right")
        fig.text(0.095, y, label, fontsize=13, color=INK_2)

    section_label(fig, 0.58, 0.80, "After: the report")
    kpi_tiles(fig, 0.58, 0.65, 0.38, 0.13, [
        (f"${summary['Net sales (shipped and delivered)']:,.0f}", "net sales, 6 months"),
        (f"{int(summary['Values fixed in the export']):,}", "values fixed"),
        (f"{int(summary['Rows removed (blank, footer, duplicate, unreadable)']):,}", "junk rows removed"),
    ])
    ax = fig.add_axes([0.61, 0.13, 0.35, 0.44])
    style_axis(ax)
    bars = ax.bar(months["Month"].str.slice(0, 3), months["Sales"], color=BLUE, width=0.6)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v / 1000:,.0f}k" if v else "$0"))
    best = months["Sales"].idxmax()
    for i in (best, len(months) - 1):
        bar = bars[i]
        ax.text(bar.get_x() + bar.get_width() / 2, bar.get_height(), f"${months['Sales'][i]:,.0f}",
                ha="center", va="bottom", fontsize=13, color=INK_2)
    ax.set_title("Net sales by month", loc="left", fontsize=15, color=INK_2, pad=12)
    fig.text(0.58, 0.075, "Also in the workbook: by product, by state, clean data, fixes log.",
             fontsize=13, color=INK_2)

    footer(fig, "Python  ·  pandas  ·  openpyxl", "All data is made-up sample data")
    save(fig, "sales-report.png")


# ---------- ads waste finder ----------

OUTCOME_STYLE = [  # outcome in the report, label on the chart, colour
    ("No sales", "No sales: block", STATUS["critical"]),
    ("Too expensive", "Too expensive: lower bid", STATUS["serious"]),
    ("Keep watching", "Keep watching", MUTED),
    ("Profitable", "Profitable", STATUS["good"]),
]


def ads_image():
    demo = HERE / "ads-waste-finder"
    report = demo / "output" / "ads_report.xlsx"
    raw = pd.read_csv(demo / "sample_data" / "search_term_report.csv", dtype=str)
    raw.columns = [c.strip() for c in raw.columns]
    negatives = pd.read_excel(report, sheet_name="Negatives to add")
    blocked = set(negatives["Search Term"])
    sample = raw.head(10)
    rows = [[r["Date"], shorten(r["Customer Search Term"], 30), r["Clicks"], f"${float(r['Spend']):.2f}",
             r["7 Day Total Orders (#)"]] for _, r in sample.iterrows()]
    highlight = {(n, 1) for n, term in enumerate(sample["Customer Search Term"]) if term in blocked}

    summary = load_summary(report)
    outcomes = pd.read_excel(report, sheet_name="Summary", header=None, usecols="A:D",
                             skiprows=17, nrows=4, names=["Outcome", "Terms", "Spend", "Share"])
    worst = negatives.groupby("Search Term")["Spend"].sum().sort_values(ascending=False).head(5)

    fig = canvas("Amazon ads waste finder",
                 "Finds the searches that cost money and never sell, from Amazon's own report")

    section_label(fig, 0.04, 0.80, f"Before: Amazon's search term report ({len(raw):,} daily rows)")
    raw_table(fig, [0.04, 0.42, 0.48, 0.37], ["Date", "Customer search term", "Clicks", "Spend", "Orders"],
              rows, highlight, [0.22, 0.43, 0.11, 0.13, 0.11], font_size=12)
    fig.text(0.04, 0.38, "Shaded = a search it tells you to block.", fontsize=13, color=INK_2)
    fig.text(0.04, 0.30, "Biggest wasters this month (no sales, all campaigns)", fontsize=15,
             fontweight="bold", color=INK_2)
    for n, (term, spend) in enumerate(worst.items()):
        y = 0.265 - n * 0.026
        fig.text(0.04, y, term, fontsize=13, color=INK_2)
        fig.text(0.40, y, f"${spend:,.2f}", fontsize=13, fontweight="bold", ha="right")

    section_label(fig, 0.56, 0.80, "After: the report")
    kpi_tiles(fig, 0.56, 0.65, 0.40, 0.13, [
        (f"${summary['Wasted spend per 30 days']:,.0f}", "wasted per month"),
        (f"{summary['Share of ad spend wasted']:.0%}", "of ad spend"),
        (f"{int(summary['Negatives to add'])}", "searches to block"),
    ])
    fig.text(0.56, 0.595, "Where the ad money went (30 days)", fontsize=15, color=INK_2)
    ax = fig.add_axes([0.73, 0.26, 0.21, 0.30])
    style_axis(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRID, linewidth=1)
    ax.spines["bottom"].set_visible(False)
    styles = {name: (label, colour) for name, label, colour in OUTCOME_STYLE}
    positions = list(range(len(outcomes)))[::-1]
    ax.barh(positions, outcomes["Spend"], color=[styles[o][1] for o in outcomes["Outcome"]], height=0.6)
    ax.set_yticks(positions, [styles[o][0] for o in outcomes["Outcome"]], fontsize=13, color=INK_2)
    ax.set_xlim(0, outcomes["Spend"].max() * 1.45)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:,.0f}"))
    for y, spend, share in zip(positions, outcomes["Spend"], outcomes["Share"]):
        ax.text(spend + outcomes["Spend"].max() * 0.03, y, f"${spend:,.0f}  ({share:.0%})",
                va="center", fontsize=12, color=INK_2)
    fig.text(0.56, 0.17, f"Also in the workbook: {int(summary['Bids to lower'])} bids to lower, each with a "
                         f"suggested bid,\nand {int(summary['Winners to add'])} winning searches to add as "
                         "exact keywords.", fontsize=13, color=INK_2, linespacing=1.5)

    footer(fig, "Python  ·  pandas  ·  openpyxl", "All data is made-up sample data")
    save(fig, "ads-waste-finder.png")


# ---------- price tracker ----------

def price_image():
    demo = HERE / "price-tracker"
    report = demo / "output" / "price_report.xlsx"
    newest = sorted((demo / "snapshots").glob("*.csv"))[-1]
    snapshot = pd.read_csv(newest, encoding="utf-8-sig")
    rows = [[shorten(r.title, 30), f"£{r.price:.2f}", "★" * int(r.rating), r.stock]
            for r in snapshot.head(10).itertuples()]
    summary = load_summary(report)
    changes = pd.read_excel(report, sheet_name="Price changes")
    top = changes.reindex(changes["Change %"].abs().sort_values(ascending=False).index).head(10)

    fig = canvas("Competitor price tracker",
                 "Checks competitor prices and stock every day, and reports what changed")

    section_label(fig, 0.04, 0.80, "What it collects on every run")
    raw_table(fig, [0.04, 0.42, 0.47, 0.37], ["Product", "Price", "Stars", "Stock"], rows, set(),
              [0.50, 0.15, 0.16, 0.19], font_size=12)
    fig.text(0.04, 0.38, "Each run is saved as a dated snapshot, so a price history builds up.",
             fontsize=13, color=INK_2)
    fig.text(0.04, 0.30, "Report sheets", fontsize=15, fontweight="bold", color=INK_2)
    for n, line in enumerate(["Summary with the biggest moves", "Price changes (old, new, % change)",
                              "New products and products gone", "Stock changes",
                              "All products, with links"]):
        fig.text(0.05, 0.265 - n * 0.022, f"·  {line}", fontsize=13, color=INK_2)
    fig.text(0.04, 0.09, "Demo note: the practice site used here never changes its prices, so the\n"
                         "earlier snapshot was generated to show what the report looks like.",
             fontsize=12, color=MUTED, linespacing=1.5)

    section_label(fig, 0.56, 0.80, "This week's report")
    kpi_tiles(fig, 0.56, 0.65, 0.40, 0.13, [
        (f"{int(summary['Price drops']) + int(summary['Price rises'])}", "price changes"),
        (f"{int(summary['New products'])}", "new products"),
        (f"{int(summary['Back in stock'])}", "back in stock"),
    ])
    ax = fig.add_axes([0.73, 0.13, 0.23, 0.44])
    style_axis(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRID, linewidth=1)
    ax.spines["bottom"].set_visible(False)
    positions = range(len(top))[::-1]
    colours = [BLUE if v < 0 else RED for v in top["Change %"]]
    ax.barh(list(positions), top["Change %"], color=colours, height=0.62)
    ax.axvline(0, color=BASELINE, linewidth=1.2)
    ax.set_yticks(list(positions), [shorten(t, 26) for t in top["Product"]], fontsize=12, color=INK_2)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:+.0%}" if v else "0%"))
    limit = top["Change %"].abs().max() * 1.7
    ax.set_xlim(-limit, limit)
    for y, value in zip(positions, top["Change %"]):
        ax.text(value + (0.012 if value > 0 else -0.012), y, f"{value:+.1%}", va="center",
                ha="left" if value > 0 else "right", fontsize=11, color=INK_2)
    fig.text(0.56, 0.595, "Biggest price moves", fontsize=15, color=INK_2)
    fig.patches.extend([
        FancyBboxPatch((0.80, 0.598), 0.012, 0.012, boxstyle="square,pad=0", facecolor=BLUE,
                       edgecolor="none", transform=fig.transFigure, figure=fig),
        FancyBboxPatch((0.88, 0.598), 0.012, 0.012, boxstyle="square,pad=0", facecolor=RED,
                       edgecolor="none", transform=fig.transFigure, figure=fig),
    ])
    fig.text(0.816, 0.597, "price drop", fontsize=12, color=INK_2)
    fig.text(0.896, 0.597, "price rise", fontsize=12, color=INK_2)

    footer(fig, "Python  ·  requests  ·  BeautifulSoup  ·  pandas  ·  openpyxl",
           "Demo site: books.toscrape.com (built for scraping practice)")
    save(fig, "price-tracker.png")


# ---------- invoice extractor ----------

def invoices_image():
    import pdfplumber  # only this image needs it

    demo = HERE / "invoice-extractor"
    report = demo / "output" / "invoices.xlsx"
    pdfs = sorted((demo / "sample_data" / "invoices").glob("*.pdf"))
    picks = [next(p for p in pdfs if p.name.startswith(prefix)) for prefix in ("northside", "brightspark", "coastal")]
    pages = []
    for pdf_path in picks:
        with pdfplumber.open(pdf_path) as pdf:
            pages.append(pdf.pages[0].crop((0, 0, pdf.pages[0].width, pdf.pages[0].height * 0.62))
                         .to_image(resolution=110).original)

    summary = load_summary(report)
    invoices = pd.read_excel(report, sheet_name="Invoices")
    checks = pd.read_excel(report, sheet_name="Checks")
    supplier_rows = pd.read_excel(report, sheet_name="Summary", header=None, skiprows=12, usecols="A:E",
                                  names=["Supplier", "Invoices", "Ex GST", "GST", "Total"]).dropna()
    supplier_rows["Total"] = supplier_rows["Total"].astype(float)

    fig = canvas("PDF invoices to Excel", "A folder of supplier invoices in, one checked spreadsheet out")

    section_label(fig, 0.04, 0.80, f"Before: {len(invoices)} PDF invoices, 3 different layouts")
    for n, page in enumerate(pages):
        width = 0.30
        height = width * fig.get_figwidth() / fig.get_figheight() * page.height / page.width
        ax = fig.add_axes([0.04 + n * 0.09, 0.415 - n * 0.06, width, height])
        ax.imshow(page)
        ax.set_xticks([])
        ax.set_yticks([])
        for spine in ax.spines.values():
            spine.set_color(BASELINE)
            spine.set_linewidth(1.2)
    fig.text(0.04, 0.215, "Ruled tables, \"500 x item @ $0.09\" lines, and hours x rate,\n"
                          "with dates written three different ways.", fontsize=13, color=INK_2, linespacing=1.5)
    fig.text(0.04, 0.165, "Each supplier's layout is read as it comes. Nothing is retyped.",
             fontsize=13, color=INK_2)

    section_label(fig, 0.56, 0.80, "After: one checked workbook")
    kpi_tiles(fig, 0.56, 0.65, 0.40, 0.13, [
        (f"{int(summary['Invoices (duplicates removed)'])}", "invoices"),
        (f"${summary['Total spend, inc GST']:,.0f}", "total spend"),
        (f"{int(summary['Invoices to check'])}", "flagged to check"),
    ])
    fig.text(0.56, 0.595, "Flagged for a person to check", fontsize=15, color=INK_2)
    for n, row in enumerate(checks.itertuples()):
        y = 0.555 - n * 0.045
        fig.patches.append(FancyBboxPatch((0.56, y + 0.004), 0.008, 0.018, boxstyle="square,pad=0",
                                          facecolor=STATUS["critical"], edgecolor="none",
                                          transform=fig.transFigure, figure=fig))
        fig.text(0.575, y + 0.013, f"{row[3]}", fontsize=13, fontweight="bold", va="center")
        text = shorten(plain_check(row[4]), 58).replace("$", r"\$")  # two $ would start matplotlib maths
        fig.text(0.655, y + 0.013, text, fontsize=13, color=INK_2, va="center")

    ax = fig.add_axes([0.70, 0.13, 0.24, 0.24])
    style_axis(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRID, linewidth=1)
    ax.spines["bottom"].set_visible(False)
    positions = list(range(len(supplier_rows)))[::-1]
    ax.barh(positions, supplier_rows["Total"], color=BLUE, height=0.55)
    ax.set_yticks(positions, list(supplier_rows["Supplier"]), fontsize=13, color=INK_2)
    ax.set_xlim(0, supplier_rows["Total"].max() * 1.35)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v / 1000:,.0f}k" if v else "$0"))
    for y, total in zip(positions, supplier_rows["Total"]):
        ax.text(total + supplier_rows["Total"].max() * 0.03, y, f"${total:,.0f}", va="center", fontsize=12,
                color=INK_2)
    fig.text(0.56, 0.39, "Spend by supplier, GST included", fontsize=15, color=INK_2)
    fig.text(0.56, 0.075, "Also in the workbook: every line item, due dates, and what's due in the next 14 days.",
             fontsize=13, color=INK_2)

    footer(fig, "Python  ·  pdfplumber  ·  openpyxl", "All invoices are made-up sample data")
    save(fig, "invoice-extractor.png")


# ---------- contact list cleaner ----------

SOURCE_LABELS = {"shop_customers.csv": "Shop", "event_signups.xlsx": "Sign-ups", "old_crm_export.csv": "Old CRM"}


def lists_image():
    demo = HERE / "list-cleaner"
    report = demo / "output" / "clean_contacts.xlsx"
    merged = pd.read_excel(report, sheet_name="Merged Rows", dtype=str).fillna("")
    clean_list = pd.read_excel(report, sheet_name="Clean List", dtype=str).fillna("")
    summary = load_summary(report)

    wanted = ["James Wilson", "Jennifer Anderson", "Noah Anderson", "Oliver Kelly", "Elizabeth Anderson"]
    before = merged[merged["Person"].isin(wanted)].copy()
    before["order"] = before["Person"].map(wanted.index)
    before = before.sort_values(["order", "File"])
    rows = [[SOURCE_LABELS.get(r["File"], r["File"]), shorten(r["Name as written"], 20),
             shorten(r["Email as written"], 30) or "(blank)", r["Phone as written"] or "(blank)"]
            for _, r in before.iterrows()]
    highlight = set()
    for n, (_, r) in enumerate(before.iterrows()):
        name = r["Name as written"]
        if name != name.strip() or "," in name or name.islower() or name.isupper() or "  " in name:
            highlight.add((n, 1))
        email = r["Email as written"]
        if not email or email != email.strip().lower() or " " in email or not email.split("@")[-1].count("."):
            highlight.add((n, 2))
        if not re.fullmatch(r"0[2378] \d{4} \d{4}|04\d\d \d{3} \d{3}", r["Phone as written"]):
            highlight.add((n, 3))  # not yet in the one standard format

    after = clean_list[(clean_list["First Name"] + " " + clean_list["Last Name"]).isin(wanted)
                       & (clean_list["Rows merged"].astype(int) > 1)].copy()
    after["order"] = (after["First Name"] + " " + after["Last Name"]).map(wanted.index)
    after = after.sort_values("order")
    after_rows = [[f"{r['First Name']} {r['Last Name']}", shorten(r["Email"], 25) or "(none)", r["Phone"],
                   r["Marketing OK"]] for _, r in after.iterrows()]

    fig = canvas("Contact list cleaner", "Three messy customer lists in, one clean list out, duplicates merged")

    section_label(fig, 0.04, 0.80, "Before: the same people in three lists")
    raw_table(fig, [0.04, 0.37, 0.50, 0.42], ["List", "Name as written", "Email as written", "Phone as written"],
              rows, highlight, [0.13, 0.25, 0.37, 0.25], font_size=11)
    fig.text(0.04, 0.335, "Shaded = written differently, broken or missing.", fontsize=13, color=INK_2)
    fig.text(0.04, 0.27, "What it fixed in the sample", fontsize=15, fontweight="bold", color=INK_2)
    fixed = [(int(summary["Duplicate rows merged"]), "duplicate rows merged"),
             (int(summary["Phone numbers put in one format"]), "phone numbers put in one format"),
             (int(summary["Emails tidied (case, spaces, mailto:)"]), "emails tidied (case, spaces, mailto:)"),
             (int(summary["Values to check (bad email or phone)"]), "bad emails or phones listed to check"),
             (int(summary["Blank rows removed"]), "blank rows removed")]
    for n, (count, label) in enumerate(fixed):
        y = 0.235 - n * 0.026
        fig.text(0.085, y, f"{count:,}", fontsize=13, fontweight="bold", ha="right")
        fig.text(0.095, y, label, fontsize=13, color=INK_2)

    section_label(fig, 0.58, 0.80, "After: one clean list")
    kpi_tiles(fig, 0.58, 0.65, 0.38, 0.13, [
        (f"{int(summary['People in the clean list'])}", "people, from 197 rows"),
        (f"{int(summary['Duplicate rows merged'])}", "duplicates merged"),
        (f"{int(summary['Said no or unsubscribed'])}", "opted out, kept out"),
    ])
    fig.text(0.58, 0.595, "The same five people, merged", fontsize=15, color=INK_2)
    raw_table(fig, [0.58, 0.38, 0.38, 0.20], ["Name", "Email", "Phone", "Marketing"], after_rows, set(),
              [0.29, 0.36, 0.21, 0.14], font_size=11)
    fig.text(0.58, 0.30, "Merges are made on the same email or phone. A similar name in the\n"
                         "same suburb only counts when nothing contradicts it, and is shaded\n"
                         "in the workbook for a person to check.", fontsize=13, color=INK_2, linespacing=1.5)
    fig.text(0.58, 0.235, "One \"no\" or \"unsubscribed\" in any list means no marketing.",
             fontsize=13, color=INK_2)
    fig.text(0.58, 0.195, "An email that can't be fixed is left out and listed, never guessed.",
             fontsize=13, color=INK_2)

    footer(fig, "Python  ·  openpyxl", "All data is made up (example.com emails, fictional phone ranges)")
    save(fig, "list-cleaner.png")


# ---------- customer insights ----------

def customers_image():
    demo = HERE / "customer-insights"
    report = demo / "output" / "customer_insights.xlsx"
    raw = pd.read_csv(demo / "sample_data" / "online_retail_sample.csv", dtype=str, keep_default_na=False)
    no_id = raw["CustomerID"] == ""
    copy = raw.duplicated(keep="first")
    postage = (raw["StockCode"] == "POST") & ~no_id
    returned = raw["InvoiceNo"].str.startswith("C")
    big_cancelled = raw.index[raw["CustomerID"] == "12346"]  # 74,215 jars ordered and cancelled the same day
    first_copy = raw.index[copy][0]
    twin = raw.index[(raw == raw.loc[first_copy]).all(axis=1)][0]
    picks = list(big_cancelled) + [twin, first_copy, raw.index[postage][0], raw.index[no_id][0]]
    plain = raw.index[~(no_id | copy | postage | returned | raw.duplicated(keep=False))]
    picks = sorted(set(picks) | set(plain[:3]))
    shaded = no_id | copy | postage | returned
    rows = [[r["InvoiceNo"], r["StockCode"], shorten(r["Description"].strip(), 24), r["Quantity"], r["UnitPrice"],
             r["CustomerID"] or "(blank)"] for _, r in raw.loc[picks].iterrows()]
    highlight = {(n, -1) for n, i in enumerate(picks) if shaded[i]}

    summary = load_summary(report)
    segments = pd.read_excel(report, sheet_name="Segments", nrows=9)
    log = pd.read_excel(report, sheet_name="Cleaning log", nrows=7).dropna(subset=["Step"])
    removed = dict(zip(log["Step"].astype(int), log["Lines"]))
    returns_value = -log.loc[log["Step"] == 5, "Value (£)"].iloc[0]

    fig = canvas("Customer insights (RFM)",
                 "Real order lines in: every customer scored, put in a segment, with an action for each")

    section_label(fig, 0.04, 0.80, f"Before: real order lines ({int(summary['Order lines in the file']):,} "
                                   "in the file)")
    raw_table(fig, [0.04, 0.49, 0.49, 0.30], ["Invoice", "Code", "Description", "Qty", "Price", "Customer"], rows,
              highlight, [0.14, 0.11, 0.37, 0.10, 0.11, 0.17], font_size=11)
    fig.text(0.04, 0.465, "Shaded = removed, or set aside as a return. 12346 ordered 74,215 jars and cancelled\n"
                          "them minutes later, so their spend is £0.", fontsize=13, color=INK_2, linespacing=1.5,
             va="top")
    fig.text(0.04, 0.345, "What the cleaning found in the full file", fontsize=15, fontweight="bold", color=INK_2)
    found = [(removed[1], "lines with no customer ID, left out"),
             (removed[2], "exact duplicate lines removed"),
             (removed[3], "postage, fees and other non-product lines"),
             (removed[4], "free items (price £0) removed"),
             (removed[5], f"returns, £{returns_value:,.0f} taken off spend")]
    for n, (count, label) in enumerate(found):
        y = 0.31 - n * 0.026
        fig.text(0.115, y, f"{int(count):,}", fontsize=13, fontweight="bold", ha="right")
        fig.text(0.125, y, label, fontsize=13, color=INK_2)

    section_label(fig, 0.56, 0.80, "After: every customer in a segment")
    kpi_tiles(fig, 0.56, 0.65, 0.40, 0.13, [
        (f"{int(summary['Customers analysed']):,}", "customers analysed"),
        (f"{summary['Share of revenue from Champions']:.0%}", "from Champions"),
        (f"{int(summary[AT_RISK_KEY]):,}", "customers at risk"),
    ])
    fig.text(0.56, 0.595, "Share of net revenue by segment", fontsize=15, color=INK_2)
    ax = fig.add_axes([0.69, 0.22, 0.17, 0.35])
    style_axis(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRID, linewidth=1)
    ax.spines["bottom"].set_visible(False)
    positions = list(range(len(segments)))[::-1]
    slipping = segments["Segment"].isin(["At risk", "Can't lose them"])
    ax.barh(positions, segments["Share of revenue"], height=0.62,
            color=[STATUS["serious"] if s else BLUE for s in slipping])
    ax.set_yticks(positions, list(segments["Segment"]), fontsize=13, color=INK_2)
    ax.set_xlim(0, segments["Share of revenue"].max() * 1.15)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0%}"))
    for y, share, count in zip(positions, segments["Share of revenue"], segments["Customers"]):
        ax.text(share + 0.012, y, f"{share:.1%}  ({count:,} customers)", va="center", fontsize=12, color=INK_2)
    fig.patches.append(FancyBboxPatch((0.56, 0.152), 0.012, 0.012, boxstyle="square,pad=0",
                                      facecolor=STATUS["serious"], edgecolor="none", transform=fig.transFigure,
                                      figure=fig))
    fig.text(0.577, 0.151, "Slipping away: good customers who haven't ordered for months", fontsize=12,
             color=INK_2)
    fig.text(0.56, 0.075, "Also in the workbook: every customer's scores and segment, a suggested action\n"
                          "for each segment, monthly revenue and new customers, and the cleaning log.",
             fontsize=13, color=INK_2, linespacing=1.5)

    footer(fig, "Python  ·  pandas  ·  openpyxl",
           "Data: Online Retail, UCI Machine Learning Repository (Daqing Chen, 2015), CC BY 4.0")
    save(fig, "customer-insights.png")


AT_RISK_KEY = "Customers at risk (At risk + Can't lose them)"


# ---------- certificates and letters from a spreadsheet ----------

def letter_card(fig, rect, docx_path):
    """Draw the top of a Word letter on a white page. The text is read back from the .docx file itself."""
    import textwrap

    from docx import Document

    doc = Document(docx_path)
    lines = [p.text for p in doc.paragraphs if p.text.strip()]
    org, address, date_line, name, dear, subject, congrats, fee, closing, regards, signer, role = lines
    table = [(row.cells[0].text, row.cells[1].text) for row in doc.tables[0].rows]

    width, height = rect[2] * 1600, rect[3] * 1200
    ax = fig.add_axes(rect)
    ax.set_xlim(0, width)
    ax.set_ylim(height, 0)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_facecolor("white")
    for spine in ax.spines.values():
        spine.set_color(BASELINE)
        spine.set_linewidth(1.2)
    pad, family = 20, ["Calibri", "Segoe UI", "DejaVu Sans"]

    def write(y, text, size=6.8, bold=False, color=INK, x=pad):
        ax.text(x, y, text, fontsize=size, fontweight="bold" if bold else "normal", color=color, va="top",
                fontfamily=family)

    def paragraph(y, text, chars=50):
        for line in textwrap.wrap(text, chars):
            write(y, line)
            y += 12
        return y + 6

    write(24, org, size=10, bold=True, color="#184f95")
    write(44, address, size=5.4, color=MUTED)
    ax.plot([pad, width - pad], [57, 57], color="#b7791f", linewidth=1)
    write(72, date_line)
    write(86, name, bold=True)
    write(100, dear)
    write(114, subject, bold=True)
    y = paragraph(134, congrats)
    y = paragraph(y, fee)
    for label, value in table:
        ax.add_patch(plt.Rectangle((pad, y), width - 2 * pad, 15, facecolor="white", edgecolor=BASELINE,
                                   linewidth=0.8))
        ax.add_patch(plt.Rectangle((pad, y), 82, 15, facecolor="#eaf1fb", edgecolor=BASELINE, linewidth=0.8))
        write(y + 3.5, label, size=6.2, bold=True, x=pad + 4)
        write(y + 3.5, value, size=6.2, x=pad + 88)
        y += 15
    y = paragraph(y + 12, closing)
    write(y + 2, regards)
    write(y + 30, signer, bold=True)
    write(y + 43, role, size=6.2)


def docs_image():
    import pdfplumber  # only this image needs it
    from openpyxl import load_workbook

    demo = HERE / "doc-generator"
    log = demo / "output" / "generation_log.xlsx"
    summary = load_summary(log)
    skipped = pd.read_excel(log, sheet_name="Skipped rows", dtype=str, keep_default_na=False)
    sheet = list(load_workbook(demo / "sample_data" / "course_completions.xlsx").active.iter_rows(values_only=True))
    rows, highlight = [], set()
    for n, (name, course, when, amount) in enumerate(sheet[6:14]):  # spreadsheet rows 7 to 14, one has no name
        text = [name or "", course or "", f"{when:%d/%m/%Y}" if hasattr(when, "year") else (when or ""),
                f"${amount:,.2f}" if amount is not None else ""]
        rows.append(text)
        highlight |= {(n, c) for c, value in enumerate(text) if value == ""}

    made, passed = int(summary["Certificates made (PDF)"]), int(summary["Files that passed the re-open check"])
    total_files = passed + int(summary["Files that failed the re-open check"])

    fig = canvas("Certificates and letters from a spreadsheet",
                 "One row per person in, a PDF certificate and a Word letter out, and a log of everything")

    section_label(fig, 0.04, 0.80, f"Before: a spreadsheet of {int(summary['Rows read'])} rows")
    raw_table(fig, [0.04, 0.50, 0.38, 0.27], ["Name", "Course", "Date", "Amount"], rows, highlight,
              [0.27, 0.43, 0.16, 0.14], font_size=10)
    fig.text(0.04, 0.475, "Rows 7 to 14 of the sheet. Shaded cells are missing.\n"
                          "Dates can be real Excel dates or typed text.", fontsize=13, color=INK_2,
             linespacing=1.5, va="top")
    fig.text(0.04, 0.38, "Rows skipped, and why", fontsize=15, color=INK_2)
    for n, row in enumerate(skipped.itertuples()):
        y = 0.34 - n * 0.045
        fig.patches.append(FancyBboxPatch((0.04, y + 0.004), 0.008, 0.018, boxstyle="square,pad=0",
                                          facecolor=STATUS["critical"], edgecolor="none",
                                          transform=fig.transFigure, figure=fig))
        fig.text(0.055, y + 0.013, f"Row {row[1]}", fontsize=13, fontweight="bold", va="center")
        fig.text(0.125, y + 0.013, row[6], fontsize=13, color=INK_2, va="center")
    fig.text(0.04, 0.17, "Every file is opened again after it is written, to check the\n"
                         f"name, course and certificate number are on it: {passed} of {total_files} passed.",
             fontsize=13, color=INK_2, linespacing=1.5)

    section_label(fig, 0.46, 0.80, f"After: {made} certificates and {made} letters")
    kpi_tiles(fig, 0.46, 0.665, 0.50, 0.115, [
        (f"{made}", "PDF certificates"),
        (f"{int(summary['Letters made (Word)'])}", "Word letters"),
        (f"{int(summary['Rows skipped (data missing or unreadable)'])}", "rows skipped"),
    ])
    certificate = demo / "output" / "certificates" / "certificate_005_Zoe_OBrien.pdf"
    with pdfplumber.open(certificate) as pdf:
        page = pdf.pages[0].to_image(resolution=140).original
    cert_w = 0.30
    cert_h = cert_w * fig.get_figwidth() / fig.get_figheight() * page.height / page.width
    ax = fig.add_axes([0.46, 0.635 - cert_h, cert_w, cert_h])
    ax.imshow(page)
    ax.set_xticks([])
    ax.set_yticks([])
    for spine in ax.spines.values():
        spine.set_color(BASELINE)
        spine.set_linewidth(1.2)
    letter_card(fig, [0.775, 0.285, 0.185, 0.35], demo / "output" / "letters" / "letter_005_Zoe_OBrien.docx")
    fig.text(0.46, 0.635 - cert_h - 0.04, "PDF certificate, A4 landscape", fontsize=13, color=INK_2)
    fig.text(0.775, 0.25, "Word letter, A4", fontsize=13, color=INK_2)
    fig.text(0.46, 0.13, "Also in the log: every file name, each row's check result,\nand fees by course.",
             fontsize=13, color=INK_2, linespacing=1.5)

    footer(fig, "Python  ·  reportlab  ·  python-docx  ·  openpyxl", "All names and data are made up")
    save(fig, "doc-generator.png")


# ---------- SQL reports ----------

GREEN_SQL = "#2f7d32"
SQL_TOKENS = re.compile(
    r"(--.*$)|(\"[^\"]*\")|('[^']*')|"
    r"\b(?:WITH|AS|SELECT|FROM|JOIN|ON|GROUP BY|HAVING|ORDER BY|DESC|LIMIT|WHERE|AND|OVER|CAST|INTEGER|DISTINCT|"
    r"ROWS BETWEEN|UNBOUNDED PRECEDING|CURRENT ROW)\b|"
    r"\b(?:COUNT|SUM|MIN|MAX|ROUND|julianday|LAG|RANK)(?=\()")


def sql_image():
    demo = HERE / "sql-reports"
    report = demo / "output" / "sql_report.xlsx"
    summary = load_summary(report)
    months = pd.read_excel(report, sheet_name="Sales by month")
    products = pd.read_excel(report, sheet_name="Top products")
    with sqlite3.connect(demo / "output" / "store.db") as db:
        orders_loaded = db.execute("SELECT COUNT(*) FROM orders").fetchone()[0]
    repeat, buyers = int(summary["Repeat customers"]), int(summary["Customers who bought"])

    sql = (demo / "queries" / "03_repeat_customers.sql").read_text(encoding="utf-8")
    code = [line.rstrip() for line in sql.splitlines() if not line.startswith("--")]

    fig = canvas("SQL reports in Excel",
                 "Orders in a SQLite database, SQL queries for the questions, one formatted workbook out")

    section_label(fig, 0.04, 0.80, f"Before: a question, asked in SQL ({orders_loaded} orders in the database)")
    width, height = 0.49, 0.46
    ax = fig.add_axes([0.04, 0.775 - height, width, height])
    ax.set_xlim(0, width * 1600)
    ax.set_ylim(height * 1200, 0)
    ax.set_xticks([])
    ax.set_yticks([])
    ax.set_facecolor(TILE)
    for spine in ax.spines.values():
        spine.set_color(GRID)
    mono, size = ["Consolas", "DejaVu Sans Mono"], 9.5
    probe = fig.text(0, 0, "M" * 40, fontsize=size, fontfamily=mono)
    advance = probe.get_window_extent(fig.canvas.get_renderer()).width / 40  # one character, in pixels
    probe.remove()
    for n, line in enumerate(code):
        y, position, parts = 28 + n * 21.5, 0, []
        for match in SQL_TOKENS.finditer(line):
            if match.start() > position:
                parts.append((position, line[position:match.start()], INK, "normal"))
            if match.group(1):  # a comment
                parts.append((match.start(), match.group(0), MUTED, "normal"))
            elif match.group(2) or match.group(3):  # a quoted name or text
                parts.append((match.start(), match.group(0), GREEN_SQL, "normal"))
            else:  # a keyword or function
                parts.append((match.start(), match.group(0), BLUE, "bold"))
            position = match.end()
        if position < len(line):
            parts.append((position, line[position:], INK, "normal"))
        for start, text, color, weight in parts:
            ax.text(18 + start * advance, y, text, fontsize=size, fontfamily=mono, va="center", ha="left",
                    color=color, fontweight=weight)
    fig.text(0.04, 0.275, "One .sql file for each report. Anyone can re-run them in any SQLite tool,\n"
                          "and the workbook's last sheet shows the SQL behind every report.",
             fontsize=13, color=INK_2, linespacing=1.5, va="top")
    fig.text(0.04, 0.16, "Checked: net sales, orders and the monthly, product and state totals\n"
                         "match the Python sales report to the cent.",
             fontsize=13, color=INK_2, linespacing=1.5, va="top")

    section_label(fig, 0.56, 0.80, "After: one Excel workbook")
    kpi_tiles(fig, 0.56, 0.665, 0.40, 0.115, [
        (f"${summary['Net sales']:,.0f}", "net sales"),
        (f"{int(summary['Orders'])}", "orders"),
        (f"{repeat}", "repeat customers"),
    ])
    fig.text(0.56, 0.615, "Net sales by month", fontsize=15, color=INK_2)
    ax = fig.add_axes([0.58, 0.395, 0.38, 0.195])
    style_axis(ax)
    ax.bar(range(len(months)), months["Sales"], color=BLUE, width=0.62)
    ax.set_xticks(range(len(months)), [m.split()[0] for m in months["Month"]], fontsize=13, color=INK_2)
    ax.set_ylim(0, months["Sales"].max() * 1.22)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v / 1000:,.0f}k" if v else "$0"))
    for x, sales in enumerate(months["Sales"]):
        ax.text(x, sales + months["Sales"].max() * 0.03, f"${sales:,.0f}", ha="center", fontsize=12, color=INK_2)

    fig.text(0.56, 0.345, "Net sales by product", fontsize=15, color=INK_2)
    ax = fig.add_axes([0.72, 0.115, 0.24, 0.205])
    style_axis(ax)
    ax.grid(axis="y", visible=False)
    ax.grid(axis="x", color=GRID, linewidth=1)
    ax.spines["bottom"].set_visible(False)
    positions = list(range(len(products)))[::-1]
    ax.barh(positions, products["Sales"], color=BLUE, height=0.62)
    ax.set_yticks(positions, [shorten(p, 30) for p in products["Product"]], fontsize=11, color=INK_2)
    ax.set_xlim(0, products["Sales"].max() * 1.3)
    ax.xaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v / 1000:,.0f}k" if v else "$0"))
    for y, sales in zip(positions, products["Sales"]):
        ax.text(sales + products["Sales"].max() * 0.02, y, f"${sales:,.0f}", va="center", fontsize=11, color=INK_2)
    fig.text(0.56, 0.075, f"Also in the workbook: {repeat} repeat customers and totals by state.",
             fontsize=13, color=INK_2)

    footer(fig, "Python  ·  SQLite  ·  openpyxl", "All data is made up")
    save(fig, "sql-reports.png")


def plain_check(text):
    """Short versions of the Checks sheet wording for the image."""
    if text.startswith("GST charged"):
        return text.replace("GST charged is", "GST is").replace(", but 10% of the subtotal is", "; 10% would be")
    if text.startswith("Subtotal"):
        return "total is " + text.split("(")[-1].rstrip(")").lstrip("+") + " more than subtotal + GST"
    if text.startswith("Same invoice"):
        return "same invoice saved twice, counted once"
    return text


def shorten(text, limit):
    text = str(text)
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def load_summary(report):
    ws = pd.read_excel(report, sheet_name="Summary", header=None, usecols="A:B", nrows=14)
    return {str(k).strip(): v for k, v in zip(ws[0], ws[1]) if isinstance(k, str) and pd.notna(v)}


def save(fig, name):
    OUT.mkdir(exist_ok=True)
    path = OUT / name
    fig.savefig(path, facecolor=SURFACE)
    plt.close(fig)
    print(f"Wrote {path}")


IMAGES = {"sales": sales_image, "ads": ads_image, "prices": price_image, "invoices": invoices_image,
          "lists": lists_image, "customers": customers_image, "docs": docs_image, "sql": sql_image}

if __name__ == "__main__":
    wanted = sys.argv[1:] or list(IMAGES)
    for key in wanted:
        IMAGES[key]()
