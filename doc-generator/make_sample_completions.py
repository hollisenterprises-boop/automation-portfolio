"""Write the sample spreadsheet for the document generator demo.

    python make_sample_completions.py

Makes sample_data/course_completions.xlsx: 40 course completions (name, course,
date, amount paid) plus one blank row, the way a training provider's roster
tends to look. Three rows are incomplete on purpose (no name, no amount, no
course or date), four dates are typed as text in different formats instead of
being real Excel dates, and a few names have accents, apostrophes or hyphens.

All names and data are made up. The random seed is fixed, so every run gives
the same file.
"""

from __future__ import annotations

import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

OUT = Path(__file__).parent / "sample_data" / "course_completions.xlsx"
SEED = 2026

FIRST = ["Olivia", "Jack", "Charlotte", "Noah", "Amelia", "William", "Isla", "Oliver", "Mia", "Thomas",
         "Grace", "James", "Chloe", "Lucas", "Ava", "Henry", "Sophie", "Liam", "Ruby", "Samuel",
         "Ella", "Benjamin", "Harper", "Daniel", "Matilda", "Joshua", "Priya", "Arjun", "Mei", "Tariq",
         "Aroha", "Wei", "Fatima", "Kai", "Leila", "Marcus", "Nadia", "Declan", "Imogen", "Rohan"]
LAST = ["Smith", "Nguyen", "Brown", "Wilson", "Taylor", "Johnson", "White", "Martin", "Anderson", "Thompson",
        "Walker", "Harris", "Lee", "Ryan", "Robinson", "Kelly", "King", "Davies", "Wright", "Mitchell",
        "Campbell", "Scott", "Murphy", "Hughes", "Edwards", "Clarke", "Price", "Turner", "Nair", "Patel",
        "Singh", "Chen", "Tran", "Hassan", "Williams", "Cooper", "Bennett", "Morgan", "Reid", "Stewart"]
COURSES = [("Excel Essentials", 395.00), ("Bookkeeping Basics", 450.00), ("Introduction to Python", 595.00),
           ("Social Media for Small Business", 295.00), ("Public Speaking Workshop", 350.00)]
FIRST_DAY, LAST_DAY = date(2026, 2, 2), date(2026, 9, 25)

# Row positions (0 = the first person) that get a special case.
SPECIAL_NAMES = {3: "Zoë O'Brien", 10: "José Álvarez", 16: "Mary-Jane Whitfield-Clarke",
                 24: "Alexandria Wilhelmina Featherstonehaugh-Montgomery"}
TEXT_DATES = {5: "%d/%m/%Y", 12: "%Y-%m-%d", 20: "{d} %b %Y", 28: "{d} %B %Y"}  # typed as text, not Excel dates
NO_NAME, NO_AMOUNT, NO_COURSE_OR_DATE = 7, 18, 32
BLANK_ROW_AFTER = 26  # a blank spacer row after this person

HEADER = ["Name", "Course", "Date", "Amount"]


def make_people(count: int = 40) -> list[dict]:
    """Chronological roster: name, course, date completed and amount paid."""
    rng = random.Random(SEED)  # a new generator each time, so every call gives the same people
    names = [f"{first} {last}" for first, last in zip(rng.sample(FIRST, count), rng.sample(LAST, count))]
    span = (LAST_DAY - FIRST_DAY).days
    days = sorted(rng.sample(range(span + 1), count))
    people = []
    for position, (name, offset) in enumerate(zip(names, days)):
        course, fee = rng.choice(COURSES)
        early_bird = rng.random() < 0.2
        people.append({
            "name": SPECIAL_NAMES.get(position, name),
            "course": course,
            "date": FIRST_DAY + timedelta(days=offset),
            "amount": round(fee * 0.9, 2) if early_bird else fee,
        })
    return people


def sheet_rows(people: list[dict]) -> list[list | None]:
    """The sheet's rows after the header; None is a blank row."""
    rows: list[list | None] = []
    for position, person in enumerate(people):
        name, course, when, amount = person["name"], person["course"], person["date"], person["amount"]
        if position in TEXT_DATES:
            pattern = TEXT_DATES[position].replace("{d}", str(when.day))
            when = when.strftime(pattern)
        else:
            when = datetime.combine(when, datetime.min.time())
        if position == NO_NAME:
            name = None
        if position == NO_AMOUNT:
            amount = None
        if position == NO_COURSE_OR_DATE:
            course = when = None
        rows.append([name, course, when, amount])
        if position == BLANK_ROW_AFTER:
            rows.append(None)
    return rows


def write_sample(path: Path = OUT) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Completions"
    ws.append(HEADER)
    for cell in ws[1]:
        cell.fill, cell.font = PatternFill("solid", fgColor="184F95"), Font(bold=True, color="FFFFFF")
    for row in sheet_rows(make_people()):
        ws.append(row or [None] * len(HEADER))
    for row in ws.iter_rows(min_row=2):
        row[2].number_format = "dd/mm/yyyy"
        row[3].number_format = '"$"#,##0.00'
    for letter, width in zip("ABCD", (34, 32, 14, 12)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    wb.properties.creator = "make_sample_completions.py"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    print(f"Wrote {write_sample()}")
