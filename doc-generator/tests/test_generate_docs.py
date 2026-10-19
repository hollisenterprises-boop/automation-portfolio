import csv
import sys
from datetime import date, datetime
from pathlib import Path

import pdfplumber
import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import generate_docs as gd  # noqa: E402
import make_sample_completions as sample  # noqa: E402

ISSUED = date(2026, 10, 19)
HEADER = ["Name", "Course", "Date", "Amount"]


def write_csv(path, rows, header=HEADER):
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(header)
        writer.writerows(rows)
    return path


@pytest.fixture(scope="module")
def sheet(tmp_path_factory):
    return sample.write_sample(tmp_path_factory.mktemp("sheet") / "course_completions.xlsx")


@pytest.fixture(scope="module")
def run(sheet, tmp_path_factory):
    out = tmp_path_factory.mktemp("out")
    return gd.generate(sheet, out, issued=ISSUED), out


def made_row(result, number):
    return next(m for m in result.made if m.row.number == number)


# ---------- counts and file names ----------

def test_reads_every_row_with_data(run):
    result, _ = run
    assert len(result.rows) == 40
    assert result.blank_rows == 1  # the spacer row in the middle of the sheet
    assert len(result.made) == 37
    assert len(result.skipped) == 3


def test_makes_one_certificate_and_one_letter_per_complete_row(run):
    result, out = run
    assert len(list((out / "certificates").glob("*.pdf"))) == 37
    assert len(list((out / "letters").glob("*.docx"))) == 37
    assert {m.row.number for m in result.made}.isdisjoint({r.number for r in result.skipped})


def test_file_names(run):
    result, out = run
    zoe = made_row(result, 5)
    assert (zoe.certificate, zoe.letter) == ("certificates/certificate_005_Zoe_OBrien.pdf",
                                             "letters/letter_005_Zoe_OBrien.docx")
    for made in result.made:
        assert (out / made.certificate).is_file() and (out / made.letter).is_file()
    names = [m.certificate for m in result.made]
    assert len(set(names)) == len(names)


@pytest.mark.parametrize("text, want", [
    ("Priya Nair", "Priya_Nair"),
    ("Zoë O'Brien", "Zoe_OBrien"),
    ("Mary-Jane Whitfield-Clarke", "Mary_Jane_Whitfield_Clarke"),
    ("José Álvarez", "Jose_Alvarez"),
    ("  ../../evil name  ", "evil_name"),
    ("张伟", "unnamed"),
    ("A" * 80, "A" * 40),
])
def test_slug(text, want):
    assert gd.slug(text) == want


# ---------- missing and unreadable data ----------

def test_incomplete_rows_are_skipped_with_the_reason(run):
    result, out = run
    reasons = {r.number: r.problems for r in result.skipped}
    assert reasons == {9: ["Missing: name"], 20: ["Missing: amount"], 35: ["Missing: course, date"]}
    assert not list((out / "certificates").glob("certificate_009_*"))
    assert not list((out / "letters").glob("letter_020_*"))


@pytest.mark.parametrize("blank, expected", [
    (0, "Missing: name"), (1, "Missing: course"), (2, "Missing: date"), (3, "Missing: amount"),
])
def test_each_missing_field_is_named(tmp_path, blank, expected):
    record = ["Priya Nair", "Excel Essentials", "12/03/2026", "395"]
    record[blank] = ""
    rows, _ = gd.read_rows(write_csv(tmp_path / "in.csv", [record]))
    assert rows[0].problems == [expected]


def test_unreadable_and_odd_values_are_flagged(tmp_path):
    rows, _ = gd.read_rows(write_csv(tmp_path / "in.csv", [
        ["Priya Nair", "Excel Essentials", "TBC", "395"],
        ["Tom Lee", "Excel Essentials", "12/03/2026", "n/a"],
        ["Ann Ray", "Excel Essentials", "12/03/2026", "-5"],
        ["Nguyễn An", "Excel Essentials", "12/03/2026", "395"],
        ["B" * 90, "Excel Essentials", "12/03/2026", "395"],
        ["Cy Dee", "Very long course name " * 8, "12/03/2026", "395"],
    ]))
    assert [r.problems for r in rows] == [
        ['Could not read the date "TBC"'],
        ['Could not read the amount "n/a"'],
        ["The amount is negative (-5)"],
        ["The name has characters the certificate font can't print: ễ"],
        ["The name is too long to fit on the certificate"],
        ["The course name is too long to fit on the certificate"],
    ]


@pytest.mark.parametrize("value, want", [
    ("12/03/2026", date(2026, 3, 12)),  # day first
    ("2026-03-12", date(2026, 3, 12)),
    ("12 Mar 2026", date(2026, 3, 12)),
    ("12 March 2026", date(2026, 3, 12)),
    (datetime(2026, 3, 12, 9, 30), date(2026, 3, 12)),
    ("next Tuesday", None),
    ("31/02/2026", None),
])
def test_parse_date(value, want):
    assert gd.parse_date(value) == want


@pytest.mark.parametrize("value, want", [
    ("$1,250.50", 1250.50), ("AUD 395", 395.0), (395, 395.0), (395.456, 395.46),
    ("n/a", None), ("nan", None), (True, None),
])
def test_parse_amount(value, want):
    assert gd.parse_amount(value) == want


# ---------- what is in the files ----------

def test_every_file_passes_its_check(run):
    result, _ = run
    assert [m.row.number for m in result.made if m.check != "OK"] == []


def test_certificate_contents(run):
    result, out = run
    for number, name in ((5, "Zoë O'Brien"), (12, "José Álvarez")):
        made = made_row(result, number)
        with pdfplumber.open(out / made.certificate) as pdf:
            assert len(pdf.pages) == 1
            page = pdf.pages[0]
            assert page.width > page.height  # landscape
            text = gd.squash(page.extract_text())
        assert gd.squash(name) in text
        assert made.row.course in text
        assert made.number in text
        assert f"{made.row.completed.day} {made.row.completed:%B %Y}" in text


def test_the_longest_name_is_shrunk_to_fit_on_one_line(run):
    result, out = run
    made = next(m for m in result.made if len(m.row.name) > 45)
    assert gd.name_size(made.row.name) < 30
    with pdfplumber.open(out / made.certificate) as pdf:
        assert gd.squash(made.row.name) in gd.squash(pdf.pages[0].extract_text())


def test_letter_contents(run):
    result, out = run
    made = made_row(result, 5)
    doc = Document(out / made.letter)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Dear Zoë," in text
    assert "Public Speaking Workshop" in text
    assert "19 October 2026" in text
    table = {row.cells[0].text: row.cells[1].text for row in doc.tables[0].rows}
    assert table == {"Course": made.row.course, "Completed": "25 March 2026",
                     "Certificate number": "C-2026-005", "Fee paid": "$350.00"}
    assert round(doc.sections[0].page_width.cm, 1) == 21.0  # A4, not US Letter


def test_the_check_notices_a_file_that_belongs_to_someone_else(run):
    result, out = run
    one, other = made_row(result, 2), made_row(result, 3)
    assert gd.check_certificate(out / one.certificate, other.row, other.number) == "The name is not on the certificate"
    assert gd.check_letter(out / one.letter, other.row, other.number) == "The name is not in the letter"
    assert gd.check_certificate(out / "letters" / "missing.pdf", one.row, one.number).startswith("The certificate would not open")


# ---------- the log and the totals ----------

def test_log_workbook(run):
    result, out = run
    wb = load_workbook(out / "generation_log.xlsx")
    assert wb.sheetnames == ["Summary", "Files made", "Skipped rows"]
    assert wb["Files made"].max_row == 1 + 37
    assert wb["Skipped rows"].max_row == 1 + 3
    skipped = [[c.value for c in row] for row in wb["Skipped rows"].iter_rows(min_row=2)]
    assert skipped[0][0] == 9 and skipped[0][-1] == "Missing: name"
    assert [row[0] for row in skipped] == [9, 20, 35]
    summary = {wb["Summary"].cell(row=r, column=1).value: wb["Summary"].cell(row=r, column=2).value
               for r in range(4, 13)}
    assert summary["Rows read"] == 40
    assert summary["Certificates made (PDF)"] == 37
    assert summary["Letters made (Word)"] == 37
    assert summary["Rows skipped (data missing or unreadable)"] == 3
    assert summary["Files that passed the re-open check"] == 74
    assert summary["Files that failed the re-open check"] == 0
    assert summary["Blank rows ignored"] == 1


def test_totals_match_the_spreadsheet(run, sheet):
    result, out = run
    source = [[c.value for c in row] for row in load_workbook(sheet).active.iter_rows(min_row=2)]
    complete = [r for r in source if all(v not in (None, "") for v in r)]
    partial = [r for r in source if any(v not in (None, "") for v in r) and r not in complete]
    assert len(complete) == 37 and len(partial) == 3

    ws = load_workbook(out / "generation_log.xlsx")["Summary"]
    labels = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(4, 13)}
    assert labels["Fees on the documents made"] == pytest.approx(sum(r[3] for r in complete))
    assert labels["Fees on skipped rows (nothing made)"] == pytest.approx(sum(r[3] for r in partial if r[3]))

    by_course = [[c.value for c in row] for row in ws.iter_rows(min_row=18, max_col=3) if row[0].value]
    assert sum(r[1] for r in by_course) == 37
    assert sum(r[2] for r in by_course) == pytest.approx(labels["Fees on the documents made"])
    assert {r[0] for r in by_course} == {r[1] for r in complete}


# ---------- reading different spreadsheets ----------

def test_csv_and_xlsx_are_read_the_same(sheet, tmp_path):
    rows = [[c.value for c in row] for row in load_workbook(sheet).active.iter_rows(min_row=2)]
    text_rows = [[gd.show(v) for v in row] for row in rows]
    from_csv, blanks = gd.read_rows(write_csv(tmp_path / "same.csv", text_rows))
    from_xlsx, _ = gd.read_rows(sheet)
    assert blanks == 1
    assert [(r.number, r.name, r.course, r.completed, r.amount, r.problems) for r in from_csv] == \
           [(r.number, r.name, r.course, r.completed, r.amount, r.problems) for r in from_xlsx]


def test_column_names_can_differ_and_blank_rows_at_the_end_are_ignored(tmp_path):
    path = tmp_path / "other.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["Notes", "  Full   Name ", "Course name", "Date completed", "Fee paid"])
    ws.append(["x", "Priya Nair", "Excel Essentials", datetime(2026, 3, 12), 395])
    for _ in range(5):
        ws.append([None] * 5)
    wb.save(path)
    rows, blanks = gd.read_rows(path)
    assert blanks == 0
    assert [(r.number, r.name, r.completed, r.amount, r.ok) for r in rows] == [
        (2, "Priya Nair", date(2026, 3, 12), 395.0, True)]


def test_a_missing_column_is_explained(tmp_path):
    path = write_csv(tmp_path / "bad.csv", [["Priya", "Excel", "12/03/2026"]], header=["Name", "Course", "Date"])
    with pytest.raises(ValueError, match="no amount column"):
        gd.read_rows(path)


# ---------- running it ----------

def test_a_failure_on_one_row_does_not_stop_the_others(tmp_path, monkeypatch):
    path = write_csv(tmp_path / "in.csv", [
        ["Priya Nair", "Excel Essentials", "12/03/2026", "395"],
        ["Tom Lee", "Bookkeeping Basics", "13/03/2026", "450"],
    ])
    real = gd.draw_letter

    def flaky(file, row, number, settings, issued):
        if row.name == "Tom Lee":
            raise OSError("disk full")
        return real(file, row, number, settings, issued)

    monkeypatch.setattr(gd, "draw_letter", flaky)
    result = gd.generate(path, tmp_path / "out", issued=ISSUED)
    assert [m.row.name for m in result.made] == ["Priya Nair"]
    assert result.skipped[0].problems == ["Could not make the files (disk full)"]
    assert [p.name for p in (tmp_path / "out" / "certificates").iterdir()] == ["certificate_002_Priya_Nair.pdf"]


def test_files_from_an_earlier_run_are_noticed(tmp_path):
    path = write_csv(tmp_path / "in.csv", [["Priya Nair", "Excel Essentials", "12/03/2026", "395"]])
    first = gd.generate(path, tmp_path / "out", issued=ISSUED)
    assert first.leftover == 0
    (tmp_path / "out" / "certificates" / "old_certificate.pdf").write_bytes(b"%PDF-1.4")
    assert gd.generate(path, tmp_path / "out", issued=ISSUED).leftover == 1  # the same names are simply replaced


def test_running_twice_gives_the_same_files_and_numbers(sheet, tmp_path):
    rows = [[gd.show(c.value) for c in row] for row in load_workbook(sheet).active.iter_rows(min_row=2, max_row=9)]
    path = write_csv(tmp_path / "first_rows.csv", rows)  # includes the row with no name
    one = gd.generate(path, tmp_path / "one", issued=ISSUED)
    two = gd.generate(path, tmp_path / "two", issued=ISSUED)
    assert len(one.made) == 7 and len(one.skipped) == 1
    assert [(m.certificate, m.letter, m.number) for m in one.made] == [(m.certificate, m.letter, m.number) for m in two.made]
    assert [r.problems for r in one.skipped] == [r.problems for r in two.skipped]


def test_command_line_and_settings(tmp_path, capsys):
    path = write_csv(tmp_path / "in.csv", [["Priya Nair", "Excel Essentials", "12/03/2026", "395"],
                                           ["", "Excel Essentials", "13/03/2026", "395"]])
    gd.main([str(path), "-o", str(tmp_path / "out"), "--as-of", "2026-10-19",
             "--org", "Acme Training", "--signer", "Pat Lee", "--note", ""])
    printed = capsys.readouterr().out
    assert "Read 2 rows from in.csv: made 1 certificates and 1 letters, skipped 1." in printed
    assert "Row 3 skipped: Missing: name" in printed
    with pdfplumber.open(tmp_path / "out" / "certificates" / "certificate_002_Priya_Nair.pdf") as pdf:
        text = gd.squash(pdf.pages[0].extract_text())
    assert "ACME TRAINING" in text and "Pat Lee" in text and "Sample document" not in text
    letter = Document(tmp_path / "out" / "letters" / "letter_002_Priya_Nair.docx")
    assert letter.paragraphs[0].text == "Acme Training"
    assert letter.sections[0].footer.paragraphs[0].text == ""


# ---------- the committed sample and output ----------

DEMO = Path(gd.__file__).parent


def sheet_values(path):
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in load_workbook(path)}


def test_the_committed_sample_is_what_the_script_writes(tmp_path):
    fresh = sample.write_sample(tmp_path / "fresh.xlsx")
    assert sheet_values(fresh) == sheet_values(DEMO / "sample_data" / "course_completions.xlsx")


def test_the_committed_log_is_what_a_run_writes(run):
    _, out = run
    assert sheet_values(out / "generation_log.xlsx") == sheet_values(DEMO / "output" / "generation_log.xlsx")


def test_the_committed_examples_are_real_files():
    certificates = sorted((DEMO / "output" / "certificates").glob("*.pdf"))
    letters = sorted((DEMO / "output" / "letters").glob("*.docx"))
    assert len(certificates) == len(letters) == 3
    for path in certificates:
        with pdfplumber.open(path) as pdf:
            assert len(pdf.pages) == 1
    for path in letters:
        assert Document(path).tables
