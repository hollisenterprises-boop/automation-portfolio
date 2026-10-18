"""Tests for the customer insights demo. Run from the demo folder with:  python -m pytest tests"""

import io
import sys
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

DEMO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(DEMO))
import analyse as an  # noqa: E402

sys.path.insert(0, str(DEMO.parent))
from make_images import load_summary  # noqa: E402

SAMPLE = DEMO / "sample_data" / "online_retail_sample.csv"
OUTPUT = DEMO / "output" / "customer_insights.xlsx"

# A few lines in the dataset's own layout, one or two for each rule.
TINY = """\
InvoiceNo,StockCode,Description,Quantity,InvoiceDate,UnitPrice,CustomerID,Country
536365,85123A,WHITE HANGING HEART T-LIGHT HOLDER,6,2010-12-01 08:26,2.55,17850,United Kingdom
536365,85123A,WHITE HANGING HEART T-LIGHT HOLDER,6,2010-12-01 08:26,2.55,17850,United Kingdom
536365,POST,POSTAGE,1,2010-12-01 08:26,18.00,17850,United Kingdom
536366,22633,HAND WARMER UNION JACK,6,2010-12-01 08:28,1.85,,United Kingdom
536366,22633,HAND WARMER UNION JACK,6,2010-12-01 08:28,1.85,,United Kingdom
536367,22423,REGENCY CAKESTAND 3 TIER,1,2010-12-05 14:02,0,12647,Germany
536368,22423,REGENCY CAKESTAND 3 TIER,2,2010-12-05 14:05,12.75,12647,Germany
C536369,85123A,WHITE HANGING HEART T-LIGHT HOLDER,-2,2010-12-08 10:00,2.55,17850,United Kingdom
536370,M,Manual,1,2010-12-06 09:00,10.00,12647,Germany
C536371,D,Discount,-1,2010-12-06 09:05,5.00,12647,Germany
C536372,C2,CARRIAGE,-1,2010-12-06 09:06,50.00,12647,Germany
536373,21730,GLASS STAR FROSTED T-LIGHT HOLDER,4,2010-12-02 10:00,4.25,13047,United Kingdom
C536374,21730,GLASS STAR FROSTED T-LIGHT HOLDER,-4,2010-12-02 11:00,4.25,13047,United Kingdom
C536375,22086,PAPER CHAIN KIT 50'S CHRISTMAS,-10,2010-12-04 12:00,2.95,14000,United Kingdom
536376,22086,damaged,-5,2010-12-04 12:00,0,,United Kingdom
536377,BANK CHARGES,Bank Charges,1,2010-12-07 10:00,15.00,15000,United Kingdom
536378,gift_0001_20,Dotcomgiftshop Gift Voucher £20.00,1,2010-12-07 11:00,17.02,16000,United Kingdom
"""


def tiny():
    return an.analyse(an.prepare(pd.read_csv(io.StringIO(TINY), dtype=str)))


def log_counts(result, column="Lines"):
    log = result["log"].set_index("Rule")
    rules = [an.NO_CUSTOMER, an.DUPLICATE, an.NOT_A_PRODUCT, an.ZERO, an.RETURN]
    return [int(log.loc[rule, column]) for rule in rules]


@pytest.fixture(scope="module")
def sample():
    return an.analyse(an.load(SAMPLE))


# ---------- cleaning ----------

def test_each_cleaning_rule_and_its_count():
    result = tiny()
    # no customer, duplicate, not a product, free item, returns
    assert log_counts(result) == [3, 1, 4, 1, 4]
    # Before any other rule: both copies of the no-customer line are duplicates, and
    # the no-customer write-off (quantity -5, price 0) breaks the zero rule too.
    assert log_counts(result, "In the whole file") == [3, 2, 4, 2, 5]
    assert len(result["sales"]) == 4


def test_every_line_is_counted_once():
    result = tiny()
    log = result["log"].set_index("Rule")
    removed = sum(log_counts(result)[:4])
    assert removed + len(result["returns"]) + len(result["sales"]) == log.loc["Order lines in the file", "Lines"]
    assert log.loc["Sale lines used", "Lines"] == len(result["sales"])


def test_not_a_product_codes_from_the_file():
    codes = pd.Series(["POST", "DOT", "C2", "BANK CHARGES", "AMAZONFEE", "CRUK", "S", "B", "PADS",
                       "gift_0001_20", "post", "M", "m", "D", "85123A", "85123a", "DCGS0070"])
    flags = an.not_goods(codes).tolist()
    assert flags == [True] * 11 + [False] * 6  # Manual, Discount and real products stay in


def test_manual_and_discount_lines_stay_in_spend():
    customers = tiny()["customers"]
    # 2 cake stands at £12.75 + a £10 manual line - a £5 discount; the free stand and the £50 carriage are out
    assert customers.loc[12647, "Sales"] == pytest.approx(35.50)
    assert customers.loc[12647, "Returns"] == pytest.approx(5.00)
    assert customers.loc[12647, "Net spend"] == pytest.approx(30.50)


def test_returns_cut_spend_but_are_not_orders():
    result = tiny()
    c = result["customers"].loc[17850]
    assert c["Sales"] == pytest.approx(15.30)   # 6 x £2.55, the duplicate and postage left out
    assert c["Returns"] == pytest.approx(5.10)  # 2 returned
    assert c["Net spend"] == pytest.approx(10.20)
    assert c["Orders"] == 1  # the return invoice is not an order
    assert c["Last order"] == pd.Timestamp("2010-12-01 08:26")  # a later return doesn't make them recent
    assert result["as_of"] == pd.Timestamp("2010-12-09")  # the day after the last invoice
    assert c["Days since last order"] == 8
    assert result["customers"].loc[12647, "Orders"] == 2  # the free-item invoice is gone


def test_customers_with_no_net_spend_are_left_out():
    result = tiny()
    assert set(result["customers"].index) == {17850, 12647}
    assert result["left_out"]["Why"].to_dict() == {
        13047: "Returned as much as they bought, or more",
        14000: "Returns only, no orders in the period",
        15000: "Nothing left after cleaning (only free items or charges)",
        16000: "Nothing left after cleaning (only free items or charges)",
    }


def test_sample_cleaning_counts(sample):
    raw = pd.read_csv(SAMPLE, dtype=str)
    assert sample["lines"] == len(raw) == 4105
    assert log_counts(sample) == [13, 6, 43, 2, 123]
    assert log_counts(sample, "In the whole file") == [13, 6, 47, 11, 127]
    assert len(sample["sales"]) == 3918
    assert raw["CustomerID"].isna().sum() == 13
    assert raw["InvoiceNo"].str.startswith("C").sum() == 127


# ---------- scoring ----------

def test_scores_split_customers_into_fifths():
    values = pd.Series(range(1, 11), dtype=float)
    edges = an.cut_points(values)
    assert edges == pytest.approx([2.8, 4.6, 6.4, 8.2])
    assert an.score(values, edges).tolist() == [1, 1, 2, 2, 3, 3, 4, 4, 5, 5]
    assert an.score(values, edges, reverse=True).tolist() == [5, 5, 4, 4, 3, 3, 2, 2, 1, 1]


def test_value_on_a_cut_point_goes_in_the_lower_band():
    edges = [18.0, 26.0, 34.0, 42.0]
    values = pd.Series([18.0, 18.01, 42.0, 42.01, 0.0, 1000.0])
    assert an.score(values, edges).tolist() == [1, 2, 4, 5, 1, 5]
    assert an.score(values, edges, reverse=True).tolist() == [5, 4, 2, 1, 5, 1]  # recency: fewer days is better


def test_equal_values_always_get_the_same_score():
    orders = pd.Series([1, 1, 1, 1, 1, 1, 2, 2, 3, 9])
    edges = an.cut_points(orders)
    assert edges == pytest.approx([1.0, 1.0, 1.4, 2.2])
    assert an.score(orders, edges).tolist() == [1, 1, 1, 1, 1, 1, 4, 4, 5, 5]


def test_value_score_rounds_up():
    pairs = [(1, 1), (1, 2), (2, 2), (3, 4), (4, 5), (5, 5)]
    assert [an.value_score(f, m) for f, m in pairs] == [1, 2, 2, 4, 5, 5]


def test_every_score_combination_has_one_segment():
    combos = {(r, f, m): an.segment_of(r, f, m) for r in range(1, 6) for f in range(1, 6) for m in range(1, 6)}
    assert len(combos) == 125
    assert set(combos.values()) == set(an.SEGMENTS)
    expected = {(5, 5, 5): "Champions", (4, 5, 4): "Champions", (3, 5, 5): "Loyal", (5, 4, 3): "Loyal",
                (5, 1, 1): "New", (4, 1, 3): "Potential loyalists", (3, 3, 3): "Need attention",
                (2, 3, 3): "At risk", (1, 5, 5): "Can't lose them", (2, 1, 2): "Hibernating", (1, 1, 1): "Lost",
                (1, 2, 2): "Lost"}
    assert {k: combos[k] for k in expected} == expected


def test_every_customer_gets_exactly_one_segment(sample):
    customers, left_out = sample["customers"], sample["left_out"]
    assert customers.index.is_unique
    assert customers["Segment"].isin(list(an.SEGMENTS)).all()
    assert (customers["RFM"] == customers["R"].astype(str) + customers["F"].astype(str)
            + customers["M"].astype(str)).all()
    assert customers["Segment"].tolist() == [an.segment_of(*s) for s in zip(customers["R"], customers["F"],
                                                                             customers["M"])]
    everyone = set(pd.read_csv(SAMPLE)["CustomerID"].dropna().astype(int))
    assert set(customers.index).isdisjoint(left_out.index)
    assert set(customers.index) | set(left_out.index) == everyone
    assert (customers["Net spend"] > 0).all() and (left_out["Net spend"] <= 0).all()


def test_segment_revenue_adds_up_to_total_net_revenue(sample):
    segments = sample["segments"]
    # every sale minus every return, less the customers left out (whose net spend is zero or less)
    total = (sample["sales"]["Value"].sum() + sample["returns"]["Value"].sum()
             - sample["left_out"]["Net spend"].sum())
    assert segments["Net revenue"].sum() == pytest.approx(sample["net_revenue"], abs=0.005)
    assert sample["net_revenue"] == pytest.approx(total, abs=0.005)
    assert segments["Customers"].sum() == len(sample["customers"])
    assert segments["Share of revenue"].sum() == pytest.approx(1.0)


def test_monthly_figures_add_up(sample):
    monthly = sample["monthly"]
    assert monthly["Net revenue"].sum() == pytest.approx(sample["net_revenue"], abs=0.005)
    assert monthly["New customers"].sum() == len(sample["customers"])
    assert (monthly["Active customers"] >= monthly["New customers"]).all()
    assert monthly["Note"].iloc[0].startswith("First month in the data")


# ---------- workbook ----------

def test_workbook(sample, tmp_path):
    path = tmp_path / "insights.xlsx"
    an.build_workbook(sample, SAMPLE.name).save(path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Segments", "Customers", "Monthly", "Cleaning log"]
    assert wb["Customers"].max_row == 1 + len(sample["customers"])
    summary = load_summary(path)
    assert summary["Order lines in the file"] == 4105
    assert summary["Customers analysed"] == len(sample["customers"])
    assert summary["Net revenue (after returns)"] == pytest.approx(sample["net_revenue"], abs=0.005)
    assert summary["Share of revenue from Champions"] == pytest.approx(sample["champion_share"])


def test_committed_workbook_from_the_full_dataset():
    assert OUTPUT.stat().st_size < 2_000_000
    summary = load_summary(OUTPUT)
    for key in ("Customers analysed", "Share of revenue from Champions",
                "Customers at risk (At risk + Can't lose them)", "Order lines in the file"):
        assert key in summary  # make_covers.py reads these
    assert summary["Order lines in the file"] == 541909
    customers = pd.read_excel(OUTPUT, sheet_name="Customers")
    assert len(customers) == summary["Customers analysed"]
    assert customers["Customer ID"].is_unique
    assert customers["Net spend"].is_monotonic_decreasing
    segments = pd.read_excel(OUTPUT, sheet_name="Segments", nrows=len(an.SEGMENTS))
    assert segments["Net revenue"].sum() == pytest.approx(summary["Net revenue (after returns)"], abs=1)


def test_missing_columns_are_named():
    with pytest.raises(ValueError, match="UnitPrice, CustomerID"):
        an.prepare(pd.DataFrame(columns=["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate",
                                         "Country"]))
