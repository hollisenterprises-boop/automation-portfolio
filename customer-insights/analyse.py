"""Find a store's best customers and the ones slipping away, from its order lines.

    python analyse.py
    python analyse.py "path/to/Online Retail.xlsx" -o output/customer_insights.xlsx
    python analyse.py sample_data/online_retail_sample.csv -o sample_insights.xlsx

Reads every order line (Excel or CSV, laid out like the UCI Online Retail
file), cleans it and logs each rule with its line count. Then it scores every
customer 1 to 5 on how recently they ordered (R), how often (F) and how much
they spent after returns (M), and puts each one in a segment such as
Champions, At risk or Lost, with a suggested action. With no file given it
reads the full dataset from data/Online Retail.xlsx (see the README for the download).
"""

from __future__ import annotations

import argparse
import re
from datetime import date
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
DEFAULT_DATA = HERE / "data" / "Online Retail.xlsx"  # not in the repo: see the README
DEFAULT_OUTPUT = HERE / "output" / "customer_insights.xlsx"
SOURCE = "Online Retail dataset, UCI Machine Learning Repository (Daqing Chen, 2015), CC BY 4.0"

COLUMNS = ["InvoiceNo", "StockCode", "Description", "Quantity", "InvoiceDate", "UnitPrice",
           "CustomerID", "Country"]

# Stock codes that are charges or bookkeeping, not goods (checked against the full file).
# Manual (M) and Discount (D) lines stay in: they are money the customer paid or got
# back, often a refund of a real order.
NOT_GOODS = ["POST", "DOT", "C2", "BANK CHARGES", "AMAZONFEE", "CRUK", "S", "B", "PADS"]
GIFT_VOUCHER = re.compile(r"^gift_\d+_\d+$", re.IGNORECASE)

NO_CUSTOMER = "No customer ID"
DUPLICATE = "Exact duplicate of an earlier line"
NOT_A_PRODUCT = ("Not a product: postage and carriage (POST, DOT, C2), bank charges, Amazon fees, "
                 "charity commission (CRUK), samples (S), bad-debt adjustments (B), gift vouchers, "
                 "cushion pads at £0.001 (PADS)")
ZERO = "Sale with a quantity or price of zero or less"
RETURN = "Cancellation or return (invoice number starts with C)"

# Recency score (5 = ordered most recently) -> segment, for value scores 1 to 5.
# Value score = the average of the frequency and spend scores, rounded up.
GRID = {
    5: ["New", "Potential loyalists", "Potential loyalists", "Loyal", "Champions"],
    4: ["New", "Potential loyalists", "Potential loyalists", "Loyal", "Champions"],
    3: ["Need attention", "Need attention", "Need attention", "Loyal", "Loyal"],
    2: ["Hibernating", "Hibernating", "At risk", "At risk", "Can't lose them"],
    1: ["Lost", "Lost", "At risk", "At risk", "Can't lose them"],
}
SEGMENTS = {  # segment: (who they are, suggested action), best first
    "Champions": ("Ordered recently, order often and spend the most",
                  "Thank them and look after them first: early access to new lines, a loyalty perk. "
                  "No discounts needed."),
    "Loyal": ("Order often and spend well, though not always very recently",
              "Offer higher-value lines and bundles, and ask for reviews and referrals."),
    "Potential loyalists": ("Ordered recently, with a few orders or one bigger order",
                            "Build the habit: suggest related products and give a reason to order again soon."),
    "New": ("One small order, placed recently",
            "Welcome them: a thank-you, what else they might like, a small offer on the second order."),
    "Need attention": ("Middling customers whose last order was a while ago",
                       "Send a timely offer based on what they bought before, before they drift away."),
    "At risk": ("Used to order often or spend well, but not for a long time",
                "Win them back: a personal email, a good offer, and ask if anything went wrong."),
    "Can't lose them": ("Among the biggest spenders, but not for a long time",
                        "Contact them personally now (the owner, not a mailing list), find out why "
                        "and fix it."),
    "Hibernating": ("Low spend and few orders, a good while ago",
                    "Low-cost reminders only: new products and seasonal offers."),
    "Lost": ("Low spend and few orders, and the longest time since their last order",
             "One last win-back email, then stop spending money on them."),
}
AT_RISK = ["At risk", "Can't lose them"]

MONEY = '"£"#,##0.00'
WHOLE_MONEY = '"£"#,##0'
COUNT = "#,##0"
PERCENT = "0.0%"
ONE_DP = "0.0"
HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")
GREY = Font(color="52514E")
BAR_COLOUR = "2A78D6"


# ---------- reading ----------

def read_raw(path: Path) -> pd.DataFrame:
    """Read the order lines from Excel or CSV as they are, keeping codes and invoice numbers as text."""
    text = {"InvoiceNo": str, "StockCode": str, "Description": str, "Country": str}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return pd.read_excel(path, dtype=text)
    return pd.read_csv(path, dtype=text, encoding="utf-8-sig")


def load(path: Path) -> pd.DataFrame:
    return prepare(read_raw(path))


def prepare(raw: pd.DataFrame) -> pd.DataFrame:
    """Check the columns, set each one's type and add the line value (quantity x price)."""
    raw = raw.rename(columns=lambda c: str(c).strip())
    missing = [c for c in COLUMNS if c not in raw.columns]
    if missing:
        raise ValueError("the file is missing these columns: " + ", ".join(missing))
    lines = raw[COLUMNS].copy()
    for col in ("InvoiceNo", "StockCode", "Description", "Country"):
        lines[col] = lines[col].astype("string").str.strip()
    lines["Quantity"] = pd.to_numeric(lines["Quantity"]).astype("int64")
    lines["UnitPrice"] = pd.to_numeric(lines["UnitPrice"]).astype("float64")
    lines["InvoiceDate"] = pd.to_datetime(lines["InvoiceDate"])
    lines["CustomerID"] = pd.to_numeric(lines["CustomerID"]).astype("Int64")
    lines["Value"] = lines["Quantity"] * lines["UnitPrice"]
    return lines


# ---------- cleaning ----------

def not_goods(codes: pd.Series) -> pd.Series:
    upper = codes.fillna("").str.upper()
    return upper.isin(NOT_GOODS) | codes.fillna("").str.match(GIFT_VOUCHER)


def is_return(lines: pd.DataFrame) -> pd.Series:
    return lines["InvoiceNo"].fillna("").str.upper().str.startswith("C")


def clean(lines: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Return (sale lines, return lines, cleaning log).

    The rules run in order and each line is counted once, under the first rule it
    meets. Returns are not removed: they are kept apart and taken off spend.
    """
    returned = is_return(lines)
    rules = [
        (NO_CUSTOMER, lines["CustomerID"].isna(), "Removed: it can't be tied to a customer"),
        (DUPLICATE, lines.duplicated(subset=COLUMNS, keep="first"), "Removed: the first copy is kept"),
        (NOT_A_PRODUCT, not_goods(lines["StockCode"]), "Removed from sales and from returns"),
        (ZERO, ~returned & ((lines["Quantity"] <= 0) | (lines["UnitPrice"] <= 0)), "Removed"),
    ]
    log = [{"Step": None, "Rule": "Order lines in the file", "Lines": len(lines),
            "Value": lines["Value"].sum(), "In the whole file": None, "What happens": ""}]
    keep = pd.Series(True, index=lines.index)
    for step, (rule, mask, action) in enumerate(rules, start=1):
        hit = keep & mask
        log.append({"Step": step, "Rule": rule, "Lines": int(hit.sum()), "Value": lines.loc[hit, "Value"].sum(),
                    "In the whole file": int(mask.sum()), "What happens": action})
        keep &= ~mask
    sales, returns = lines[keep & ~returned], lines[keep & returned]
    log.append({"Step": len(rules) + 1, "Rule": RETURN, "Lines": len(returns), "Value": returns["Value"].sum(),
                "In the whole file": int(returned.sum()),
                "What happens": "Kept as returns: taken off the customer's spend, not counted as orders"})
    log.append({"Step": None, "Rule": "Sale lines used", "Lines": len(sales), "Value": sales["Value"].sum(),
                "In the whole file": None, "What happens": ""})
    return sales.reset_index(drop=True), returns.reset_index(drop=True), pd.DataFrame(log)


# ---------- RFM ----------

def main_country(sales: pd.DataFrame) -> pd.Series:
    """The country on most of a customer's lines (alphabetical on a tie)."""
    counts = sales.groupby(["CustomerID", "Country"]).size().rename("n").reset_index()
    counts = counts.sort_values(["CustomerID", "n", "Country"], ascending=[True, False, True])
    return counts.drop_duplicates("CustomerID").set_index("CustomerID")["Country"]


def customer_table(sales: pd.DataFrame, returns: pd.DataFrame, as_of: pd.Timestamp, all_ids=None
                   ) -> tuple[pd.DataFrame, pd.DataFrame]:
    """One row per customer with recency, frequency and net spend.

    Returns (customers, left out). Left out: customers whose returns equal or
    exceed what they bought in the period (net spend of zero or less), and any
    customer in all_ids with no lines left after cleaning.
    """
    g = sales.groupby("CustomerID")
    table = pd.DataFrame({
        "Country": main_country(sales),
        "First order": g["InvoiceDate"].min(),
        "Last order": g["InvoiceDate"].max(),
        "Orders": g["InvoiceNo"].nunique(),
        "Sales": g["Value"].sum(),
    })
    table["Days since last order"] = (as_of - table["Last order"].dt.normalize()).dt.days
    returned = -returns.groupby("CustomerID")["Value"].sum()
    everyone = table.index.union(returned.index)
    if all_ids is not None:
        everyone = everyone.union(pd.Index(all_ids, dtype=everyone.dtype))
    table = table.reindex(everyone)
    table["Sales"] = table["Sales"].fillna(0.0)
    table["Returns"] = returned.reindex(everyone, fill_value=0.0)
    table["Net spend"] = table["Sales"] - table["Returns"]
    table.index.name = "Customer ID"
    out = table["Net spend"] <= 0
    left_out = table.loc[out, ["Sales", "Returns", "Net spend"]].copy()
    no_orders = table.loc[out, "Orders"].isna()
    left_out["Why"] = np.select(
        [no_orders & (left_out["Returns"] > 0), no_orders],
        ["Returns only, no orders in the period", "Nothing left after cleaning (only free items or charges)"],
        default="Returned as much as they bought, or more")
    customers = table.loc[~out].copy()
    customers["Orders"] = customers["Orders"].astype("int64")
    customers["Days since last order"] = customers["Days since last order"].astype("int64")
    return customers, left_out


def cut_points(values: pd.Series) -> list[float]:
    """The 20th, 40th, 60th and 80th percentiles: the four edges between five scores."""
    return [float(v) for v in values.quantile([0.2, 0.4, 0.6, 0.8])]


def score(values: pd.Series, edges: list[float], reverse: bool = False) -> pd.Series:
    """1 to 5 by band. A value on a cut point goes in the lower band, so equal values
    always score the same. reverse=True gives 5 to the lowest values (recency)."""
    band = np.searchsorted(np.asarray(edges, dtype=float), values.to_numpy(dtype=float), side="left")
    return pd.Series((5 - band) if reverse else (band + 1), index=values.index, dtype="int64")


def value_score(f: int, m: int) -> int:
    """The average of the frequency and spend scores, rounded up."""
    return (f + m + 1) // 2


def segment_of(r: int, f: int, m: int) -> str:
    return GRID[r][value_score(f, m) - 1]


def score_customers(customers: pd.DataFrame) -> tuple[pd.DataFrame, dict]:
    """Add R, F and M scores, the RFM code and the segment; sort by spend."""
    edges = {
        "R": cut_points(customers["Days since last order"]),
        "F": cut_points(customers["Orders"]),
        "M": cut_points(customers["Net spend"]),
    }
    scored = customers.copy()
    scored["R"] = score(scored["Days since last order"], edges["R"], reverse=True)
    scored["F"] = score(scored["Orders"], edges["F"])
    scored["M"] = score(scored["Net spend"], edges["M"])
    scored["RFM"] = scored["R"].astype(str) + scored["F"].astype(str) + scored["M"].astype(str)
    scored["Segment"] = [segment_of(r, f, m) for r, f, m in zip(scored["R"], scored["F"], scored["M"])]
    scored = scored.sort_values(["Net spend", "Customer ID"], ascending=[False, True], kind="stable")
    return scored, edges


# ---------- summaries ----------

def segment_table(customers: pd.DataFrame) -> pd.DataFrame:
    total = customers["Net spend"].sum()
    g = customers.groupby("Segment")
    seg = pd.DataFrame({
        "Customers": g.size(),
        "Net revenue": g["Net spend"].sum(),
        "Average spend": g["Net spend"].mean(),
        "Average orders": g["Orders"].mean(),
        "Average days since last order": g["Days since last order"].mean(),
    }).reindex(list(SEGMENTS))
    seg["Customers"] = seg["Customers"].fillna(0).astype("int64")
    seg["Net revenue"] = seg["Net revenue"].fillna(0.0)
    seg["Share of customers"] = seg["Customers"] / len(customers)
    seg["Share of revenue"] = seg["Net revenue"] / total
    seg["Who they are"] = [SEGMENTS[s][0] for s in seg.index]
    seg["Suggested action"] = [SEGMENTS[s][1] for s in seg.index]
    seg = seg.rename_axis("Segment").reset_index()
    return seg[["Segment", "Who they are", "Customers", "Share of customers", "Net revenue", "Share of revenue",
                "Average spend", "Average orders", "Average days since last order", "Suggested action"]]


def monthly_table(sales: pd.DataFrame, returns: pd.DataFrame, customers: pd.DataFrame) -> pd.DataFrame:
    """Net revenue, orders, active and new customers by month, for the customers analysed."""
    sales = sales[sales["CustomerID"].isin(customers.index)]
    returns = returns[returns["CustomerID"].isin(customers.index)]
    by = sales.groupby(sales["InvoiceDate"].dt.to_period("M")).agg(
        Sales=("Value", "sum"), Orders=("InvoiceNo", "nunique"), Active=("CustomerID", "nunique"))
    months = by.index.union(returns["InvoiceDate"].dt.to_period("M").unique())
    by = by.reindex(months, fill_value=0)
    by["Returns"] = -returns.groupby(returns["InvoiceDate"].dt.to_period("M"))["Value"].sum().reindex(
        months, fill_value=0.0)
    by["Net revenue"] = by["Sales"] - by["Returns"]
    by["New customers"] = customers["First order"].dt.to_period("M").value_counts().reindex(months, fill_value=0)
    first_day, last_day = sales["InvoiceDate"].min(), sales["InvoiceDate"].max()
    notes = []
    for month in months:
        note = []
        if month == months[0]:
            note.append("first month in the data, so every customer counts as new")
        if month == months[0] and first_day.day > 1:
            note.append(f"from {first_day:%d %b}")
        if month == months[-1] and last_day.day < month.days_in_month:
            note.append(f"1 to {last_day.day} {last_day:%b} only")
        text = "; ".join(note)
        notes.append(text[:1].upper() + text[1:])
    table = pd.DataFrame({
        "Month": [m.strftime("%b %Y") for m in months],
        "Net revenue": by["Net revenue"].to_numpy(),
        "Orders": by["Orders"].to_numpy(),
        "Active customers": by["Active"].to_numpy(),
        "New customers": by["New customers"].to_numpy(),
        "Sales": by["Sales"].to_numpy(),
        "Returns": by["Returns"].to_numpy(),
        "Note": notes,
    })
    return table


def analyse(lines: pd.DataFrame) -> dict:
    """Run the whole analysis on prepared order lines."""
    sales, returns, log = clean(lines)
    as_of = lines["InvoiceDate"].max().normalize() + pd.Timedelta(days=1)
    customers, left_out = customer_table(sales, returns, as_of, lines["CustomerID"].dropna().unique())
    customers, edges = score_customers(customers)
    segments = segment_table(customers)
    at_risk = customers["Segment"].isin(AT_RISK)
    champions = customers["Segment"] == "Champions"
    return {
        "lines": len(lines),
        "sales": sales,
        "returns": returns,
        "log": log,
        "as_of": as_of,
        "first_day": lines["InvoiceDate"].min(),
        "last_day": lines["InvoiceDate"].max(),
        "customers": customers,
        "left_out": left_out,
        "edges": edges,
        "segments": segments,
        "monthly": monthly_table(sales, returns, customers),
        "net_revenue": customers["Net spend"].sum(),
        "champions": int(champions.sum()),
        "champion_share": customers.loc[champions, "Net spend"].sum() / customers["Net spend"].sum(),
        "at_risk": int(at_risk.sum()),
        "at_risk_revenue": customers.loc[at_risk, "Net spend"].sum(),
        "repeat_share": (customers["Orders"] >= 2).mean(),
    }


# ---------- writing Excel ----------

def cell_value(value):
    """pandas and numpy values as plain Python values openpyxl can write."""
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime().date()
    if value is None or (not isinstance(value, str) and pd.isna(value)):
        return None
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        return float(value)
    return value


def write_table(ws, df: pd.DataFrame, start_row: int = 1, formats: dict | None = None,
                widths: dict | None = None, autosize: bool = True) -> int:
    """Write df with a styled header at start_row; return the last row written."""
    formats, widths = formats or {}, widths or {}
    for col, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, record in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col, value in enumerate(record, start=1):
            value = cell_value(value)
            cell = ws.cell(row=r, column=col, value=value)
            fmt = formats.get(df.columns[col - 1])
            if fmt:
                cell.number_format = fmt
            elif isinstance(value, date):
                cell.number_format = "dd/mm/yyyy"
    for col, name in enumerate(df.columns, start=1):
        letter = get_column_letter(col)
        if name in widths:
            ws.column_dimensions[letter].width = widths[name]
        elif autosize:
            shown = [len(str(name))] + [len(f"{v:,.2f}") + 1 if isinstance(v, float) else len(str(v))
                                        for v in df.iloc[:, col - 1].head(500)]
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                     min(max(shown) + 3, 48))
    return start_row + len(df)


def money(value: float) -> float:
    return round(float(value), 2)


def build_workbook(result: dict, source_name: str) -> Workbook:
    customers, segments, log = result["customers"], result["segments"], result["log"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = "Customer segments: UK online gift shop"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = (f"Orders {result['first_day']:%d %b %Y} to {result['last_day']:%d %b %Y}, amounts in pounds. "
                f"Built from {source_name} on {date.today():%d %b %Y}. Data: {SOURCE}.")
    ws["A2"].font = GREY
    kpis = [
        ("Order lines in the file", result["lines"], COUNT),
        ("Sale lines used after cleaning", len(result["sales"]), COUNT),
        ("Customers analysed", len(customers), COUNT),
        ("Net revenue (after returns)", money(result["net_revenue"]), WHOLE_MONEY),
        ("Average spend per customer", money(result["net_revenue"] / len(customers)), MONEY),
        ("Repeat customers (2 or more orders)", result["repeat_share"], PERCENT),
        ("Champions", result["champions"], COUNT),
        ("Share of revenue from Champions", result["champion_share"], PERCENT),
        ("Customers at risk (At risk + Can't lose them)", result["at_risk"], COUNT),
        ("Revenue from customers now at risk", money(result["at_risk_revenue"]), WHOLE_MONEY),
        ("Customers left out (no net spend after returns)", len(result["left_out"]), COUNT),
    ]
    for offset, (label, value, fmt) in enumerate(kpis):
        ws.cell(row=4 + offset, column=1, value=label)
        cell = ws.cell(row=4 + offset, column=2, value=value)
        cell.number_format = fmt
        cell.font = Font(bold=True)

    top = 17
    ws.cell(row=top - 1, column=1, value="Revenue by segment").font = Font(bold=True, size=12)
    brief = segments[["Segment", "Customers", "Net revenue", "Share of revenue"]].copy()
    brief["Net revenue"] = brief["Net revenue"].map(money)
    write_table(ws, brief, start_row=top, formats={"Customers": COUNT, "Net revenue": WHOLE_MONEY,
                                                   "Share of revenue": PERCENT})
    notes = top + len(brief) + 2
    ws.cell(row=notes, column=1, value="How it decides").font = Font(bold=True, size=12)
    how = [
        f"Recency: days from the last order to {result['as_of']:%d %b %Y}, the day after the last invoice.",
        "Frequency: number of different invoices. Spend: what they bought minus what they returned.",
        "Each is scored 1 to 5 by fifths of the customers (5 = most recent, most orders, most spend).",
        "The recency score and the average of the other two place each customer in one segment "
        "(grid on the Segments sheet).",
        "Returns are taken off spend but are not orders. The Cleaning log sheet lists every rule.",
    ]
    for offset, text in enumerate(how, start=1):
        ws.cell(row=notes + offset, column=1, value=text)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 16

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Net revenue by segment"
    chart.legend = None
    chart.height, chart.width = 9, 16
    chart.add_data(Reference(ws, min_col=3, min_row=top, max_row=top + len(brief)), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=top + 1, max_row=top + len(brief)))
    chart.series[0].graphicalProperties.solidFill = BAR_COLOUR
    chart.series[0].graphicalProperties.line.solidFill = BAR_COLOUR
    chart.x_axis.scaling.orientation = "maxMin"  # best segment at the top
    chart.y_axis.crosses = "max"  # keeps the money axis at the bottom
    chart.y_axis.numFmt = WHOLE_MONEY
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.gapWidth = 50
    ws.add_chart(chart, "F4")

    seg_ws = wb.create_sheet("Segments")
    seg_out = segments.copy()
    for col in ("Net revenue", "Average spend"):
        seg_out[col] = seg_out[col].map(money)
    last = write_table(seg_ws, seg_out, formats={
        "Customers": COUNT, "Share of customers": PERCENT, "Net revenue": WHOLE_MONEY,
        "Share of revenue": PERCENT, "Average spend": MONEY, "Average orders": ONE_DP,
        "Average days since last order": "0"}, widths={"Who they are": 44, "Suggested action": 70})
    seg_ws.freeze_panes = "B2"
    for row in seg_ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    seg_ws.row_dimensions[1].height = 32

    grid_top = last + 3
    seg_ws.cell(row=grid_top, column=1, value="How each customer gets a segment").font = Font(bold=True, size=12)
    seg_ws.cell(row=grid_top + 1, column=1,
                value="Rows: recency score. Columns: value score, the average of the frequency and spend "
                      "scores, rounded up.").font = GREY
    header = ["Recency score"] + [f"Value {v}" for v in range(1, 6)]
    write_table(seg_ws, pd.DataFrame([[r] + GRID[r] for r in (5, 4, 3, 2, 1)], columns=header),
                start_row=grid_top + 2, autosize=False)
    band_top = grid_top + 10
    seg_ws.cell(row=band_top, column=1, value="What each score means here (cut at every fifth of the customers)"
                ).font = Font(bold=True, size=12)
    write_table(seg_ws, band_table(customers, result["edges"]), start_row=band_top + 1, autosize=False)
    for letter in "BCDEF":  # room for the segment names in the grid
        seg_ws.column_dimensions[letter].width = max(seg_ws.column_dimensions[letter].width or 0, 21)

    cust_ws = wb.create_sheet("Customers")
    cust_out = customers.reset_index()[[
        "Customer ID", "Country", "Segment", "RFM", "R", "F", "M", "Days since last order", "Orders",
        "Sales", "Returns", "Net spend", "First order", "Last order"]].copy()
    for col in ("Sales", "Returns", "Net spend"):
        cust_out[col] = cust_out[col].map(money)
    write_table(cust_ws, cust_out, formats={"Customer ID": "0", "Days since last order": COUNT, "Orders": COUNT,
                                            "Sales": MONEY, "Returns": MONEY, "Net spend": MONEY,
                                            "First order": "dd/mm/yyyy", "Last order": "dd/mm/yyyy"},
                widths={"Customer ID": 12, "Segment": 20, "RFM": 7, "R": 5, "F": 5, "M": 5,
                        "First order": 12, "Last order": 12})
    cust_ws.freeze_panes = "B2"
    cust_ws.auto_filter.ref = cust_ws.dimensions

    month_ws = wb.create_sheet("Monthly")
    month_out = result["monthly"].copy()
    for col in ("Net revenue", "Sales", "Returns"):
        month_out[col] = month_out[col].map(money)
    write_table(month_ws, month_out, formats={"Net revenue": WHOLE_MONEY, "Orders": COUNT,
                                              "Active customers": COUNT, "New customers": COUNT,
                                              "Sales": WHOLE_MONEY, "Returns": WHOLE_MONEY},
                widths={"Note": 60})
    month_ws.freeze_panes = "B2"

    log_ws = wb.create_sheet("Cleaning log")
    log_out = log.copy()
    log_out["Value"] = log_out["Value"].map(money)
    log_out = log_out.rename(columns={"Value": "Value (£)"})
    last = write_table(log_ws, log_out, formats={"Lines": COUNT, "Value (£)": MONEY, "In the whole file": COUNT},
                       widths={"Step": 12, "Rule": 60, "What happens": 44})
    for row in log_ws.iter_rows(min_row=2, max_row=last):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    log_ws.freeze_panes = "A2"
    note = last + 2
    log_ws.cell(row=note, column=1, value="Each line is counted once, under the first rule it meets. "
                                          "'In the whole file' counts every line that breaks the rule, "
                                          "before any other rule.").font = GREY
    left = result["left_out"].reset_index()
    log_ws.cell(row=note + 2, column=1,
                value=f"Customers left out of the segments: {len(left)} (no net spend: their returns equal or "
                      "exceed what they bought, or nothing was left after cleaning)"
                ).font = Font(bold=True, size=12)
    if len(left):
        for col in ("Sales", "Returns", "Net spend"):
            left[col] = left[col].map(money)
        write_table(log_ws, left.sort_values(["Net spend", "Customer ID"]), start_row=note + 3,
                    formats={"Customer ID": "0", "Sales": MONEY, "Returns": MONEY, "Net spend": MONEY},
                    autosize=False)
    return wb


def band_table(customers: pd.DataFrame, edges: dict) -> pd.DataFrame:
    """The range of values behind each score, in plain words."""
    days, orders = customers["Days since last order"], customers["Orders"]
    rows = [[s, whole_band(edges["R"], 5 - s, int(days.min()), " days"),
             whole_band(edges["F"], s - 1, int(orders.min())), money_band(edges["M"], s - 1)]
            for s in (5, 4, 3, 2, 1)]
    return pd.DataFrame(rows, columns=["Score", "Days since last order", "Orders", "Net spend"])


def whole_band(cuts: list[float], band: int, lowest: int, unit: str = "") -> str:
    """Band 0 is at or below the first cut point. For whole numbers, 'v <= cut' means
    'v <= floor(cut)', so the band runs from one above the previous cut to this one."""
    low = lowest if band == 0 else int(np.floor(cuts[band - 1])) + 1
    if band == 4:
        return f"{low:,} or more{unit}"
    high = int(np.floor(cuts[band]))
    if low > high:
        return "none (tied with the score below)"
    return f"{low:,}{unit}" if low == high else f"{low:,} to {high:,}{unit}"


def money_band(cuts: list[float], band: int) -> str:
    if band == 0:
        return f"up to £{cuts[0]:,.0f}"
    if band == 4:
        return f"over £{cuts[3]:,.0f}"
    return f"£{cuts[band - 1]:,.0f} to £{cuts[band]:,.0f}"


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("data", type=Path, nargs="?", default=DEFAULT_DATA,
                        help="order lines, .xlsx or .csv (default: the UCI Online Retail file)")
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    if not args.data.is_file():
        parser.error(f"file not found: {args.data}\nDownload the data first (see README.md).")

    try:
        lines = load(args.data)
    except ValueError as exc:
        raise SystemExit(f"Could not read {args.data.name}: {exc}")
    result = analyse(lines)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(result, args.data.name).save(args.output)

    print(f"{result['lines']:,} order lines -> {len(result['sales']):,} sale lines and "
          f"{len(result['returns']):,} return lines after cleaning")
    print(f"{len(result['customers']):,} customers, £{result['net_revenue']:,.2f} net revenue; "
          f"Champions {result['champions']:,} ({result['champion_share']:.1%} of revenue), "
          f"at risk {result['at_risk']:,}. Workbook: {args.output}")


if __name__ == "__main__":
    main()
