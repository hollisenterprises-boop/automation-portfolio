"""Find the Amazon ad searches that waste money, the ones to promote and the bids to lower.

    python ads_waste_finder.py sample_data/search_term_report.csv -o output/ads_report.xlsx

Reads a Sponsored Products search term report (CSV or Excel, daily rows) as
Amazon exports it, adds up each search across the days and sorts it into one
of four outcomes: no sales, too expensive, profitable or keep watching. The
Excel report lists the negatives to add, the bids to lower (with a suggested
bid) and the winning searches to add as exact keywords.
"""

from __future__ import annotations

import argparse
import math
import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.marker import DataPoint
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Field -> headers accepted for it, Amazon's own spelling first. Headers are
# compared by normalise_header(), so case, extra spaces and "(#)" don't matter.
HEADERS = {
    "Date": ["Date"],
    "Campaign": ["Campaign Name", "Campaign"],
    "Ad Group": ["Ad Group Name", "Ad Group"],
    "Targeting": ["Targeting", "Keyword"],
    "Match Type": ["Match Type"],
    "Search Term": ["Customer Search Term", "Search Term"],
    "Impressions": ["Impressions"],
    "Clicks": ["Clicks"],
    "Spend": ["Spend", "Cost"],
    "Sales": ["7 Day Total Sales", "14 Day Total Sales", "Total Sales"],
    "Orders": ["7 Day Total Orders (#)", "14 Day Total Orders (#)", "Total Orders"],
}
OPTIONAL = {"Date", "Targeting", "Match Type"}
TEXT_FIELDS = ["Campaign", "Ad Group", "Targeting", "Match Type", "Search Term"]
COUNT_FIELDS = ["Impressions", "Clicks", "Orders"]
MONEY_FIELDS = ["Spend", "Sales"]
KEY = ["Campaign", "Ad Group", "Search Term"]

NO_SALES = "No sales"
TOO_EXPENSIVE = "Too expensive"
KEEP_WATCHING = "Keep watching"
PROFITABLE = "Profitable"
OUTCOMES = [NO_SALES, TOO_EXPENSIVE, KEEP_WATCHING, PROFITABLE]
OUTCOME_COLOURS = {NO_SALES: "D03B3B", TOO_EXPENSIVE: "EC835A",
                   KEEP_WATCHING: "898781", PROFITABLE: "0CA30C"}

NEGATIVE_EXACT = "Add as negative exact"
NEGATIVE_PRODUCT = "Add as negative product target"
LOWER_BID = "Lower bid"
ADD_EXACT = "Add as exact keyword"
KEEP_RUNNING = "Keep running"
NOT_ENOUGH_DATA = "Not enough data yet"
NEAR_TARGET = "Near target, keep running"

EXPENSIVE_FACTOR = 1.5  # "too expensive" = ACOS above 1.5 x target
WINNER_ORDERS = 2       # orders before a profitable search earns its own exact keyword
BID_STEP = 0.05
MIN_BID = 0.20

ASIN_RE = re.compile(r"b0[a-z0-9]{8}", re.IGNORECASE)

# Amazon writes "Sep 01, 2026" or 2026-09-01; slashed dates are read day-first (Australian).
DATE_FORMATS = (
    "%Y-%m-%d",
    "%Y-%m-%d %H:%M:%S",
    "%b %d, %Y",
    "%B %d, %Y",
    "%d/%m/%Y",
    "%d-%m-%Y",
    "%d %b %Y",
    "%d %B %Y",
    "%Y%m%d",
)

NEGATIVE_COLUMNS = ["Campaign", "Ad Group", "Search Term", "Clicks", "Spend", "Action"]
BID_COLUMNS = ["Campaign", "Ad Group", "Targeting", "Match Type", "Search Term",
               "Clicks", "Spend", "Sales", "Orders", "CPC", "ACOS", "Suggested Bid"]
WINNER_COLUMNS = ["Campaign", "Ad Group", "Search Term", "Clicks", "Spend",
                  "Sales", "Orders", "CPC", "ACOS"]
ALL_COLUMNS = ["Campaign", "Ad Group", "Search Term", "Targeting", "Match Type",
               "Impressions", "Clicks", "Spend", "Sales", "Orders", "CPC", "ACOS",
               "Outcome", "Action", "Suggested Bid"]

MONEY = '"$"#,##0.00'
WHOLE_MONEY = '"$"#,##0'
COUNT = "#,##0"
PERCENT = "0.0%"
FORMATS = {
    "Impressions": COUNT, "Clicks": COUNT, "Orders": COUNT, "Search terms": COUNT,
    "Spend": MONEY, "Sales": MONEY, "CPC": MONEY, "Suggested Bid": MONEY,
    "ACOS": PERCENT, "Share of spend": PERCENT,
}
HEADER_FILL = PatternFill("solid", fgColor="184F95")
HEADER_FONT = Font(bold=True, color="FFFFFF")


# ---------- reading single values ----------

def normalise_header(name) -> str:
    """'7 Day Total Orders (#)' -> '7 day total orders'."""
    return " ".join(re.sub(r"[^a-z0-9]+", " ", str(name).lower()).split())


def clean_text(value) -> str:
    return "" if pd.isna(value) else " ".join(str(value).split())


def text_key(value) -> str:
    """Compare keyword and search term text ignoring case and spacing."""
    return clean_text(value).lower()


def _number(value, strip: str) -> float:
    if pd.isna(value):
        return 0.0
    text = re.sub(strip, "", str(value))
    if text in ("", "-"):
        return 0.0
    number = float(text)
    if not math.isfinite(number):
        raise ValueError(f"not a number: {value!r}")
    return number


def parse_money(value) -> float:
    """Read 1.23, $1.23, A$1.23, AU$1.23, AUD 1.23 or 1,234.56. A blank cell is 0."""
    return round(_number(value, r"(?i)aud|au?\$|\$|,|\s"), 2)


def parse_count(value) -> int:
    """Read 12, 1,234 or 12.0. A blank cell is 0."""
    return int(round(_number(value, r"[,\s]")))


def parse_date(value) -> date | None:
    text = clean_text(value)
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def is_asin(term: str) -> bool:
    """A search term like b08kx7q2lm is a product page (an ASIN), not typed words."""
    return bool(ASIN_RE.fullmatch(term.strip()))


# ---------- reading the report ----------

def load_report(path: Path) -> pd.DataFrame:
    """Read the report as text, from CSV or Excel."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        raw = pd.read_excel(path, dtype=str)
    else:
        try:
            raw = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
        except UnicodeDecodeError:  # saved from Excel as plain "CSV" on Windows
            raw = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="cp1252")
    return raw.fillna("")


def match_columns(headers) -> dict[str, str]:
    """Map each field to the report's own header; error if a required one is missing."""
    by_key: dict[str, str] = {}
    for header in headers:
        by_key.setdefault(normalise_header(header), header)
    found, missing = {}, []
    for field, spellings in HEADERS.items():
        keys = [normalise_header(s) for s in spellings]
        header = next((by_key[k] for k in keys if k in by_key), None)
        if header is not None:
            found[field] = header
        elif field not in OPTIONAL:
            missing.append(spellings[0])
    if missing:
        raise ValueError("the report is missing these columns: " + ", ".join(missing)
                         + ". Is it a Sponsored Products search term report?")
    return found


def _read_column(raw: pd.DataFrame, header: str, parse) -> list:
    values, bad = [], []
    for index, value in raw[header].items():
        try:
            values.append(parse(value))
        except ValueError:
            bad.append(f"row {index + 2} has {value!r}")  # header is row 1
    if bad:
        raise ValueError(f"column '{header.strip()}' should hold numbers but "
                         + "; ".join(bad[:3]))
    return values


def read_rows(raw: pd.DataFrame) -> pd.DataFrame:
    """Pick the columns the rules need and turn them into clean text and numbers."""
    found = match_columns(raw.columns)
    rows = pd.DataFrame(index=raw.index)
    for field in TEXT_FIELDS:
        rows[field] = raw[found[field]].map(clean_text) if field in found else ""
    rows["Date"] = raw[found["Date"]].map(parse_date) if "Date" in found else None
    for field in COUNT_FIELDS:
        rows[field] = _read_column(raw, found[field], parse_count)
    for field in MONEY_FIELDS:
        rows[field] = _read_column(raw, found[field], parse_money)
    return rows[rows["Search Term"] != ""].reset_index(drop=True)


# ---------- deciding ----------

def acos_of(spend: float, sales: float) -> float | None:
    """Ad spend / sales, rounded to 0.1% so the rules match the % shown in the report."""
    return round(spend / sales, 3) if sales > 0 else None


def classify(orders, clicks, spend, sales, target_acos=0.30, min_clicks=10, min_spend=15.0) -> str:
    """Sort one search into NO_SALES, TOO_EXPENSIVE, PROFITABLE or KEEP_WATCHING."""
    acos = acos_of(spend, sales)
    if orders == 0 or acos is None:
        seen_enough = clicks >= min_clicks or spend >= min_spend
        return NO_SALES if seen_enough else KEEP_WATCHING
    if acos > round(target_acos * EXPENSIVE_FACTOR, 6):
        return TOO_EXPENSIVE
    if acos <= round(target_acos, 6):
        return PROFITABLE
    return KEEP_WATCHING


def suggest_bid(cpc: float, acos: float, target_acos: float) -> float:
    """The bid that would have hit the target: CPC x target / ACOS,
    to the nearest 5 cents and never below 20 cents."""
    steps = math.floor(round(cpc * target_acos / acos / BID_STEP, 6) + 0.5)
    return round(max(steps * BID_STEP, MIN_BID), 2)


def exact_keywords(rows: pd.DataFrame) -> set[str]:
    """Texts that already run as an exact keyword somewhere in the report."""
    exact = rows["Match Type"].str.upper() == "EXACT"
    return {text_key(t) for t in rows.loc[exact, "Targeting"] if t}


def decide(term: dict, exact_texts: set[str], target_acos=0.30, min_clicks=10,
           min_spend=15.0) -> tuple[str, str, float | None]:
    """Return (outcome, action, suggested bid) for one search term."""
    outcome = classify(term["Orders"], term["Clicks"], term["Spend"], term["Sales"],
                       target_acos, min_clicks, min_spend)
    if outcome == NO_SALES:
        return outcome, NEGATIVE_PRODUCT if is_asin(term["Search Term"]) else NEGATIVE_EXACT, None
    if outcome == TOO_EXPENSIVE:
        return outcome, LOWER_BID, suggest_bid(term["CPC"], term["ACOS"], target_acos)
    if outcome == PROFITABLE:
        new = term["Orders"] >= WINNER_ORDERS and text_key(term["Search Term"]) not in exact_texts
        return outcome, ADD_EXACT if new else KEEP_RUNNING, None
    sold = term["Orders"] > 0 and not pd.isna(term["ACOS"])
    return outcome, NEAR_TARGET if sold else NOT_ENOUGH_DATA, None


def aggregate_terms(rows: pd.DataFrame) -> pd.DataFrame:
    """One line per campaign, ad group and search term, added up over the days."""
    sums = ["Impressions", "Clicks", "Spend", "Sales", "Orders"]
    terms = rows.groupby(KEY, sort=False, as_index=False).agg(
        **{"Targeting": ("Targeting", "first"), "Match Type": ("Match Type", "first")},
        **{field: (field, "sum") for field in sums},
    )
    terms[MONEY_FIELDS] = terms[MONEY_FIELDS].round(2)
    terms["CPC"] = [round(s / c, 2) if c else None for s, c in zip(terms["Spend"], terms["Clicks"])]
    terms["ACOS"] = [acos_of(s, v) for s, v in zip(terms["Spend"], terms["Sales"])]
    return terms


def analyse(rows: pd.DataFrame, target_acos=0.30, min_clicks=10, min_spend=15.0) -> pd.DataFrame:
    """Every search term with its totals, outcome, action and suggested bid, by spend."""
    terms = aggregate_terms(rows)
    exact = exact_keywords(rows)
    decisions = pd.DataFrame(
        [decide(t, exact, target_acos, min_clicks, min_spend) for t in terms.to_dict("records")],
        columns=["Outcome", "Action", "Suggested Bid"], index=terms.index)
    terms = pd.concat([terms, decisions], axis=1)
    return terms.sort_values(["Spend", "Search Term"], ascending=[False, True],
                             kind="stable").reset_index(drop=True)


# ---------- summarising ----------

def summarise(terms: pd.DataFrame, rows: pd.DataFrame) -> dict:
    outcome = terms["Outcome"]
    spend = round(terms["Spend"].sum(), 2)
    sales = round(terms["Sales"].sum(), 2)
    wasted = round(terms.loc[outcome == NO_SALES, "Spend"].sum(), 2)

    by_outcome = pd.DataFrame({
        "Outcome": OUTCOMES,
        "Search terms": [int((outcome == o).sum()) for o in OUTCOMES],
        "Spend": [round(terms.loc[outcome == o, "Spend"].sum(), 2) for o in OUTCOMES],
    })
    by_outcome["Share of spend"] = by_outcome["Spend"] / spend if spend else 0.0

    days = pd.to_datetime(rows["Date"]).dropna()
    first = days.min() if len(days) else None
    last = days.max() if len(days) else None
    day_count = (last - first).days + 1 if first is not None else None

    return {
        "first_day": first,
        "last_day": last,
        "days": day_count,
        "spend": spend,
        "sales": sales,
        "acos": acos_of(spend, sales),
        "wasted": wasted,
        "wasted_share": wasted / spend if spend else 0.0,
        "wasted_per_30_days": wasted / day_count * 30 if day_count else None,
        "terms": len(terms),
        "negatives": int((outcome == NO_SALES).sum()),
        "bids": int((outcome == TOO_EXPENSIVE).sum()),
        "winners": int((terms["Action"] == ADD_EXACT).sum()),
        "by_outcome": by_outcome,
    }


# ---------- writing Excel ----------

def pct(fraction: float) -> str:
    """0.3 -> '30%', 0.255 -> '25.5%'."""
    return f"{fraction * 100:g}%"


def _shown_width(value) -> int:
    if isinstance(value, float):
        return len(f"{value:,.2f}") + 1  # shown as $1,234.56
    return len(str(value))


def write_table(ws, df: pd.DataFrame, start_row: int = 1, formats: dict | None = None) -> int:
    """Write df with a styled header at start_row; return the last row written."""
    formats = formats or {}
    for col, name in enumerate(df.columns, start=1):
        cell = ws.cell(row=start_row, column=col, value=name)
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, record in enumerate(df.itertuples(index=False), start=start_row + 1):
        for col, value in enumerate(record, start=1):
            cell = ws.cell(row=r, column=col, value=None if pd.isna(value) else value)
            fmt = formats.get(df.columns[col - 1])
            if fmt:
                cell.number_format = fmt
    for col, name in enumerate(df.columns, start=1):
        width = max([len(str(name))] + [_shown_width(v) for v in df.iloc[:, col - 1].head(500)])
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0,
                                                 min(width + 3, 48))
    return start_row + len(df)


def build_workbook(terms, summary, target_acos, min_clicks, min_spend,
                   store, source_name) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"

    winners = terms[terms["Action"] == ADD_EXACT].sort_values(
        ["Sales", "Search Term"], ascending=[False, True], kind="stable")
    sheets = [
        ("Negatives to add", terms.loc[terms["Outcome"] == NO_SALES, NEGATIVE_COLUMNS]),
        ("Bids to lower", terms.loc[terms["Outcome"] == TOO_EXPENSIVE, BID_COLUMNS]),
        ("Winners to add", winners[WINNER_COLUMNS]),
        ("All search terms", terms[ALL_COLUMNS]),
    ]
    for title, table in sheets:
        sheet = wb.create_sheet(title)
        write_table(sheet, table, formats=FORMATS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    ws["A1"] = f"Amazon ads waste report: {store}"
    ws["A1"].font = Font(bold=True, size=16)
    period = ""
    if summary["days"]:
        period = (f"{summary['first_day']:%d %b %Y} to {summary['last_day']:%d %b %Y} "
                  f"({summary['days']} days). ")
    ws["A2"] = (f"{period}Target ACOS {pct(target_acos)}. "
                f"Built from {source_name} on {date.today():%d %b %Y}.")
    ws["A2"].font = Font(color="52514E")

    kpis = [
        ("Days covered", summary["days"], COUNT),
        ("Ad spend", summary["spend"], MONEY),
        ("Sales from ads", summary["sales"], MONEY),
        ("ACOS (ad spend / sales)", summary["acos"], PERCENT),
        ("Wasted spend (searches with no sales)", summary["wasted"], MONEY),
        ("Share of ad spend wasted", summary["wasted_share"], PERCENT),
        ("Wasted spend per 30 days", summary["wasted_per_30_days"], MONEY),
        ("Search terms checked", summary["terms"], COUNT),
        ("Negatives to add", summary["negatives"], COUNT),
        ("Bids to lower", summary["bids"], COUNT),
        ("Winners to add", summary["winners"], COUNT),
    ]
    kpis = [kpi for kpi in kpis if kpi[1] is not None]  # no dates or no sales: skip those lines
    for offset, (label, value, fmt) in enumerate(kpis):
        row = 4 + offset
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        cell.number_format = fmt
        cell.font = Font(bold=True)

    ws.cell(row=16, column=1, value="Ad spend by outcome").font = Font(bold=True, size=12)
    write_table(ws, summary["by_outcome"], start_row=17, formats=FORMATS)

    ws.cell(row=23, column=1, value="How each search is sorted").font = Font(bold=True, size=12)
    rules = [
        f"{NO_SALES}: no orders after {min_clicks}+ clicks or ${min_spend:,.2f}+ spend. "
        "Add it as a negative.",
        f"{TOO_EXPENSIVE}: it sells, but ACOS is above {pct(target_acos * EXPENSIVE_FACTOR)} "
        "(1.5 x target). Lower the bid to the suggested bid.",
        f"{PROFITABLE}: ACOS at or below the {pct(target_acos)} target. "
        f"With {WINNER_ORDERS}+ orders and no exact keyword yet, it is a winner to add.",
        f"{KEEP_WATCHING}: not enough clicks yet, or ACOS close to target. Look again next week.",
    ]
    for offset, text in enumerate(rules):
        ws.cell(row=24 + offset, column=1, value=text)
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 16

    chart = BarChart()
    chart.title = "Ad spend by outcome"
    chart.legend = None
    chart.height, chart.width = 7.5, 16
    chart.add_data(Reference(ws, min_col=3, min_row=17, max_row=17 + len(OUTCOMES)),
                   titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=1, min_row=18, max_row=17 + len(OUTCOMES)))
    for index, name in enumerate(OUTCOMES):
        point = DataPoint(idx=index)
        point.graphicalProperties.solidFill = OUTCOME_COLOURS[name]
        point.graphicalProperties.line.solidFill = OUTCOME_COLOURS[name]
        chart.series[0].dPt.append(point)
    chart.y_axis.numFmt = WHOLE_MONEY
    chart.x_axis.delete = False
    chart.y_axis.delete = False
    chart.gapWidth = 60
    ws.add_chart(chart, "E4")
    return wb


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("report", type=Path, help="search term report (CSV or XLSX)")
    parser.add_argument("-o", "--output", type=Path, default=Path("ads_report.xlsx"))
    parser.add_argument("--store", default="Sample Store", help="name shown on the report")
    parser.add_argument("--target-acos", type=float, default=0.30,
                        help="target ACOS as a fraction (default 0.30 = 30%%)")
    parser.add_argument("--min-clicks", type=int, default=10,
                        help="clicks with no order before a search counts as waste (default 10)")
    parser.add_argument("--min-spend", type=float, default=15.0,
                        help="spend with no order before a search counts as waste (default 15)")
    args = parser.parse_args(argv)
    if not 0 < args.target_acos <= 1:
        parser.error("--target-acos is a fraction, for example 0.30 for 30%")
    if not args.report.is_file():
        parser.error(f"file not found: {args.report}")

    try:
        rows = read_rows(load_report(args.report))
    except ValueError as exc:
        raise SystemExit(f"Could not read {args.report.name}: {exc}")
    terms = analyse(rows, args.target_acos, args.min_clicks, args.min_spend)
    summary = summarise(terms, rows)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(terms, summary, args.target_acos, args.min_clicks, args.min_spend,
                   args.store, args.report.name).save(args.output)

    period = ""
    if summary["days"]:
        period = f", {summary['first_day']:%d %b %Y} to {summary['last_day']:%d %b %Y}"
    print(f"Wasted ${summary['wasted']:,.2f} of ${summary['spend']:,.2f} ad spend "
          f"({summary['wasted_share']:.1%}) on searches with no sales{period}.")
    print(f"{summary['negatives']} negatives to add, {summary['bids']} bids to lower, "
          f"{summary['winners']} winners to add. Report: {args.output}")


if __name__ == "__main__":
    main()
